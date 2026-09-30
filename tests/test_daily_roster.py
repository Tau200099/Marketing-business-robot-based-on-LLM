"""
国庆每天的活动名单（09-29）：计划表拆成每天一份、改正生效、到哪天用哪天、同名按支局分、归组 / 通报按那天。
人名全是编的。
"""
import io
import json

from _harness import check, done, eq, run, setup

tmp = setup()

import openpyxl  # noqa: E402

from botcore import regions  # noqa: E402
from botcore.campaign import Campaign, Order  # noqa: E402
from botcore.config import ConfigStore  # noqa: E402


def plan(rows) -> bytes:
    """照真表排：第 2、3 行表头，B 列支局，E、F 列两天的活动人员。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "国庆支局营销活动计划表"
    ws["A2"], ws["B2"], ws["C2"] = "支局分类", "支局", "10月1日-10月4日国庆活动"
    ws["C3"], ws["D3"], ws["E3"], ws["F3"] = "活动类型", "目标小区", "10月1日活动人员", "10月2日活动人员"
    for i, (cls, br, d1, d2) in enumerate(rows, 4):
        ws.cell(i, 1, cls)
        ws.cell(i, 2, br)
        ws.cell(i, 5, d1)
        ws.cell(i, 6, d2)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


CITY = plan([("城区支局", "通惠", "1组：张一、李二\n2组：王三，赵四", "1组：张一 王三\n2组：李二、赵四、钱五"),
             (None, "东林", "1组：杨同、冯十", "1组：杨同、冯十\n"),
             (None, "新盛", "1组：新甲、新乙", "1组：新甲、新乙")])
RURAL = plan([("三类", "安稳", "1组：杨同、孙六 2组：周七、吴八", "1组：杨同、孙六\n2组：周七、错名")])
FIXES = {"rename": {"错名": "吴八"}, "groups": [{"day": "10.2", "branch": "安稳", "group": 2, "members": ["周七", "郑九"]}],
         "region_of": {"新盛": "农村"}}
ORDER = ["通惠", "古南", "文龙", "万盛", "东林", "打通", "赶水", "东溪", "青年", "新盛", "郭扶", "三江", "南桐", "安稳"]


def build():
    c, n1 = regions.parse_daily_plan(CITY, regions.CITY, 2026, FIXES)
    r, n2 = regions.parse_daily_plan(RURAL, regions.RURAL, 2026, FIXES)
    return regions.merge_daily([c, r], ORDER), n1 + n2


def test_parse():
    data, notes = build()
    days = data["days"]
    eq(sorted(days), ["2026-10-01", "2026-10-02"], "两天各一份")
    d1 = days["2026-10-01"]
    eq(d1["城区"]["teams"], ["通惠战队", "东林战队"], "城区战队（新盛不在城区）")
    eq(d1["农村"]["teams"], ["新盛战队", "安稳战队"], "新盛按改正算农村，排在二类的位置")
    eq(d1["城区"]["people"]["赵四"]["group"], "王三、赵四", "逗号分隔也认")
    eq(days["2026-10-02"]["城区"]["people"]["张一"]["group"], "张一、王三", "空格分隔也认；第二天换了组")
    eq(days["2026-10-02"]["城区"]["people"]["钱五"]["group"], "李二、赵四、钱五", "3 人一组")
    eq(d1["农村"]["people"]["吴八"]["group"], "周七、吴八", "「2组」和上一组挤一行也拆得开")
    eq(days["2026-10-02"]["农村"]["people"]["郑九"]["group"], "周七、郑九", "某天某组整组改正生效")
    check("错名" not in json.dumps(days, ensure_ascii=False), "错字改正生效")
    eq(notes, [], "没有同一天两个组的人")
    bad = plan([(None, "安稳", "1组：甲甲、乙乙\n2组：丙丙、甲甲", "")])
    _d, n = regions.parse_daily_plan(bad, regions.RURAL, 2026)
    check(n and "甲甲" in n[0], f"同一天两个组会提醒：{n}")


def test_by_day():
    data, _ = build()
    (tmp / regions.DAILY_FILE).write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    rs = regions.RosterSet()
    d1, d2 = rs.on("2026-10-01"), rs.on("2026-10-02 15:30:00")
    eq((d1.identify_name("王三").group, d2.identify_name("王三").group), ("王三、赵四", "张一、王三"), "同一个人按那天的组")
    eq(d1.identify_name("钱五").role, regions.UNKNOWN, "10.1 没排班的人那天不是活动人员")
    eq((d2.identify_name("钱五").role, d2.identify_name("钱五").team), (regions.ACTIVITY, "通惠战队"), "10.2 排了就是")
    eq(d1.identify_name("新甲").region, regions.RURAL, "新盛的人算农村")
    check(rs.on("2026-09-29") is rs, "没有那天的名单：用平时的名册")
    eq((d1.identify_name("杨同", "安稳").region, d1.identify_name("杨同", "东林").region), (regions.RURAL, regions.CITY),
       "同名不同人：按晒单写的 / 水印里的支局分")
    check("钱五" in rs.all_names().people, "7 天名单里的人都能按账号拼音认")
    eq(d2.identify_members("周七郑九").group, "周七、郑九", "打卡按水印小组成员也按那天的名单")


def test_placement_and_view():
    store = ConfigStore()
    rs = regions.RosterSet()
    camp = Campaign(rs.activity[regions.CITY], store.campaign_settings().targets, rosters=rs)

    def add(at, collector):
        o = Order(at=at, kind="新增", fee="199", collector=collector, team="通惠支局", msg_id=f"m{at}{collector}")
        camp.add(o)
        return o

    o2 = add("2026-10-02 10:00:00", "钱五")
    eq((o2.region, o2.role, o2.group), (regions.CITY, regions.ACTIVITY, "李二、赵四、钱五"), "10.2 的单按 10.2 的组")
    o1 = add("2026-10-01 10:00:00", "钱五")
    eq(o1.role, regions.UNKNOWN, "10.1 他没排班：不算活动人员")
    groups = [(g.label, g.members) for g in camp.view(regions.CITY, day="2026-10-02").roster.groups() if g.team == "通惠战队"]
    eq(groups, [("1组", "张一、王三"), ("2组", "李二、赵四、钱五")], "通报按那天的小组画")


def test_extra_list():
    """09-29：补充的非活动人员（渠道销售员）当揽收人算进总表；同名按支局挑；排了班那天还是活动人员。"""
    (tmp / regions.EXTRA_FILE).write_text(json.dumps({"people": [["渠道甲", "古南"], ["张同", "永新"], ["张同", "隆盛"],
                                                                  ["王三", "通惠"]]}, ensure_ascii=False),
                                          encoding="utf-8")
    rs = regions.RosterSet()
    i = rs.identify_name("渠道甲")
    eq((i.role, i.branch, i.region), (regions.OTHER, "古南", regions.CITY), "补充名单：非活动、支局按名单、城区支局算城区")
    eq((rs.identify_name("张同", "隆盛").branch, rs.identify_name("张同", "永新").branch), ("隆盛", "永新"),
       "同名不同人：按晒单写的支局挑")
    eq(rs.identify_name("张同", "隆盛").region, regions.RURAL, "农村支局算农村")
    eq(rs.on("2026-10-02").identify_name("王三").role, regions.ACTIVITY, "两边都有的人：排了班那天是活动人员")
    d1 = rs.on("2026-10-01")
    eq(d1.identify_name("杨同", "隆盛").role, regions.ACTIVITY, "隆盛没有叫杨同的：还是活动人员（东林、安稳里挑）")
    (tmp / regions.EXTRA_FILE).write_text(json.dumps({"people": [["渠道甲", "古南"], ["张同", "永新"], ["张同", "隆盛"],
                                                                  ["王三", "通惠"], ["杨同", "隆盛"]]}, ensure_ascii=False),
                                          encoding="utf-8")
    d1 = regions.RosterSet().on("2026-10-01")
    eq((d1.identify_name("杨同", "隆盛").role, d1.identify_name("杨同", "隆盛").branch), (regions.OTHER, "隆盛"),
       "活动人员和渠道销售员同名：写的是隆盛就挑隆盛那个渠道销售员")
    eq(d1.identify_name("杨同", "安稳").role, regions.ACTIVITY, "写的是安稳就是活动人员那个")
    eq(d1.identify_name("杨同").role, regions.ACTIVITY, "没写支局：活动人员优先")
    check("渠道甲" in rs.all_names().people, "补充名单里的人也能按账号拼音认")
    store = ConfigStore()
    camp = Campaign(rs.activity[regions.CITY], store.campaign_settings().targets, rosters=rs)
    o = Order(at="2026-10-02 11:00:00", kind="新增", fee="199", collector="渠道甲", team="通惠支局", msg_id="x1")
    camp.add(o)
    eq(o.branch, "古南", "渠道销售员当揽收人：总表算他名单里的支局")


run(test_parse)
run(test_by_day)
run(test_placement_and_view)
run(test_extra_list)
done()
