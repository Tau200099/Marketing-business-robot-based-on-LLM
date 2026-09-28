"""
多群第 6 块：支局长打卡点检表（功能说明第八节，用户 09-29 确认）。人脸用假的（照片 -> 认出谁、几张脸），
不用真人照片；时间、人名都是编的。
"""
import json
from datetime import datetime

from _harness import FakePlatform, check, done, eq, run, setup

tmp = setup()

from botcore.chiefs import ChiefBoard  # noqa: E402
from botcore.config import ConfigStore  # noqa: E402
from botcore.message import InboundMessage  # noqa: E402
from botcore.pipeline import Pipeline  # noqa: E402

(tmp / "chiefs.json").write_text(json.dumps({"chiefs": [
    {"class": "一类", "branch": "通惠", "name": "翁艺丹"},
    {"class": "一类", "branch": "通惠", "name": "罗江耀"},
    {"class": "二类", "branch": "郭扶", "name": "李昌洪"},
    {"class": "二类", "branch": "打通", "name": "吴佳骏"},
    {"class": "二类", "branch": "打通", "name": "吴洪霞"},
    {"class": "三类", "branch": "三江", "name": "严健齐"},
]}, ensure_ascii=False), encoding="utf-8")

# 照片 -> (认出的人, 照片里几张脸)
FAKE = {
    b"gather_0820": (["翁艺丹", "罗江耀"], 6), b"visit_1058": (["翁艺丹"], 2),
    b"visit_1539": (["罗江耀"], 2), b"visit_1542": (["罗江耀"], 3), b"alone_0815": (["李昌洪"], 1),
    b"noon_1402": (["罗江耀", "李昌洪"], 5), b"early_0720": (["吴佳骏"], 4), b"stranger": ([], 2),
    b"visit_1600": (["李昌洪"], 2),
}


class FakeFaces:
    def available(self):
        return True

    def identify(self, blob):
        names, n = FAKE[blob]
        return {"faces": n, "matches": [(x, 0.6) for x in names], "unknown": n - len(names)}


store = ConfigStore()
ADMIN = store.backup_settings().admins[0]
plat = FakePlatform()
p = Pipeline(plat, store)
board = p.chief_board
board.faces = FakeFaces()
DAY = datetime.now().strftime("%Y-%m-%d")


def photo(blob: bytes, hm: str, work: str = "上门"):
    return board.on_photo(blob, {"stamp": f"{DAY} {hm}", "work": work}, blob.decode())


def cmd(text, sender=ADMIN, chat="p2p"):
    before = len(plat.replies)
    p._process(InboundMessage(msg_id=f"c{len(plat.replies)}{text}", chat_id=sender if chat == "p2p" else chat,
                              chat_type=chat, text=text, sender_id=sender, mentions_bot=True), 0.0)
    return plat.replies[-1] if len(plat.replies) > before else ""


def row(rows, name):
    return next(r for r in rows if r.name == name)


def test_hook_installed():
    check(p.checkin.on_photo == board.submit, "打卡照片都过支局长点检（交给人脸队列，不等比对完）")
    import threading
    started = threading.Event()
    release = threading.Event()
    real = board.on_photo

    def slow(blob, wm, photo):
        started.set()
        release.wait(5)
        return real(blob, wm, photo)
    board.on_photo = slow
    import time
    t0 = time.time()
    board.submit(b"stranger", {"stamp": f"{DAY} 10:00", "work": ""}, "stranger")
    check(time.time() - t0 < 0.5, "submit 马上返回，不等人脸比对")
    started.wait(2)
    release.set()
    board.wait_idle(10)
    board.on_photo = real


def test_sightings():
    eq(photo(b"gather_0820", "08:20"), ["翁艺丹", "罗江耀"], "集合照认出两个支局长")
    photo(b"visit_1058", "10:58")
    photo(b"visit_1539", "15:39")
    photo(b"visit_1542", "15:42")
    photo(b"alone_0815", "08:15")
    photo(b"noon_1402", "14:02")
    photo(b"visit_1600", "16:00")
    eq(photo(b"early_0720", "07:20"), [], "07:30 以前的不记")
    eq(photo(b"stranger", "10:00"), [], "照片里没有支局长：什么都不存")
    eq(photo(b"visit_1058", "10:58"), ["翁艺丹"], "同一张再发：认出但不重复存")
    saved = json.loads((tmp / "chief_sightings.json").read_text(encoding="utf-8"))["sightings"]
    eq(len(saved), 9, "存了 9 条（每张照片每个认出的支局长一条）")
    check(all(set(s) == {"day", "time", "name", "faces", "work", "photo"} for s in saved), "只存人名、时间、脸数，不存照片")


def test_table_afternoon():
    rows = board.rows(DAY, datetime.now().replace(hour=15, minute=0))
    w = row(rows, "翁艺丹")
    eq((w.early, w.noon, w.visits), ("8:20", "未打卡", ["10:58"]), "翁艺丹：早会 8:20（集合照）、中午没认出 = 未打卡、第1户 10:58")
    r = row(rows, "罗江耀")
    eq((r.early, r.noon, r.visits), ("8:20", "14:02", ["15:39"]), "罗江耀：15:39、15:42 算同一户")
    li = row(rows, "李昌洪")
    eq((li.early, li.noon), ("自行打卡", "14:02"), "李昌洪：早会窗口里只有他一个人的照片 = 自行打卡")
    eq(li.visits, ["8:15", "16:00"], "自行打卡那张也算一户（和样表的李昌洪一样）")
    wj = row(rows, "吴佳骏")
    eq((wj.early, wj.noon, wj.visits), ("未打卡", "未打卡", []), "吴佳骏：07:20 那张不算，都未打卡")
    check(not any(x.red_name for x in rows), "18:00 之前不判名字标红")
    early = board.rows(DAY, datetime.now().replace(hour=8, minute=40))
    eq(row(early, "吴佳骏").early, "", "早会窗口还没过：先空着，不写未打卡")


def test_commands():
    check("只有管理员" in cmd("休息 吴洪霞", sender="YangTong"), "非管理员不能改")
    check("单聊" in cmd("休息 吴洪霞", chat="g_hq"), "群里改提示单聊")
    check(cmd("休息 吴洪霞").startswith("已标记：吴洪霞"), "管理员标休息")
    check(cmd("补记 三江 严健齐 10:37").startswith("已补记：严健齐"), "管理员补记一户（带支局）")
    check("名单里没有「张三」" in cmd("补记 张三 10:37"), "不是支局长：说清楚")
    check("要写几点" in cmd("补记 严健齐"), "补记没写时间：提示格式")
    rows = board.rows(DAY, datetime.now().replace(hour=18, minute=30))
    h = row(rows, "吴洪霞")
    eq((h.early, h.noon, h.visits, h.red_name), ("休息", "休息", ["休息"], False), "休息：三格写休息、不标红")
    eq(row(rows, "严健齐").visits, ["10:37"], "补记的一户出现在表里")


def test_red_names():
    rows = board.rows(DAY, datetime.now().replace(hour=18, minute=30))
    eq({x.name for x in rows if x.red_name}, {"翁艺丹", "吴佳骏", "严健齐"},
       "18:00 以后：14:00 以后一户都没有的标红（翁艺丹只有上午、吴佳骏没有、严健齐只有补记的 10:37）")


def test_outputs():
    now = datetime.now().replace(hour=18, minute=30)
    png = board.render_png(DAY, now)
    check(png[:8] == b"\x89PNG\r\n\x1a\n" and len(png) > 10000, f"出图（{len(png)} 字节）")
    check("日支局长打卡情况点检（7:30-18:30）" in board.title(DAY, now), board.title(DAY, now))
    import openpyxl
    wb = openpyxl.Workbook()
    board.write_sheet(wb.active, DAY, now)
    ws = wb.active
    eq([ws.cell(2, c).value for c in range(1, 7)], ["类别", "支局", "支局长", "早会", "中午集合", "第1户"], "Excel 表头和样表一样")
    eq(ws.cell(3, 3).fill.fgColor.rgb[-6:], "FF0000", "标红的名字红底")
    fp = FakePlatform()
    eq(board.send(fp, "g_hq", now), True, "发到攻坚群")


run(test_hook_installed)
run(test_sightings)
run(test_table_afternoon)
run(test_commands)
run(test_red_names)
run(test_outputs)
done()
