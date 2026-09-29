"""
出表：把 campaign.json 里的订单写成 Excel。

设计原则：**生成的表是台账，不是计算器。**
    原来那张手工表靠 VLOOKUP + 数据透视 + RANK 一层层算出来。
    机器人这边所有计算都在 Python 里做完，写进 Excel 的是【结果值】，
    不带公式。这样：
      · 打开就能看，不用等 Excel 重算
      · 不会因为插了一行就把公式引用打乱
      · 移动端、WPS、在线预览都能正常显示

「清单」的列和原表对齐，但做了三处调整（都是和你确认过的）：
    · 去掉四个从来没填过的列：用户名、地址、安装地址、区域内/外
    · 两个重名的「备注」合并成一个
    · 设备号留空 —— 那是客户手机号，机器人一律脱敏，不落盘

⚠️ 时间列只写到天（8.30），和原表格式一致。
   但 campaign.json 里存的是完整时刻，时段通报靠那个切，
   不靠这一列 —— 只到天的话上午和下午的通报会是同一份数据。
"""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

from .campaign import Campaign, GroupStat, Order, Targets, TeamStat

log = logging.getLogger("botcore.report")

# 「清单」的列定义：(表头, 取值函数, 列宽)
# 计数类的列【0 写成空】，和原表一致 —— 原表里没有 0，只有 1 或空白。
# 满屏的 0 会把真正有数的格子淹掉。
_BLANK_IF_ZERO = True


def _num(v: int):
    """计数列：0 显示为空白，和原表保持一致。"""
    return v if v or not _BLANK_IF_ZERO else None


def _day(at: str) -> str:
    """
    完整时刻 -> 原表那种「8.30」。

    解析不了就原样返回 —— 回灌历史数据时 at 本来就可能是「8.21」。
    """
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            d = datetime.strptime(at, fmt)
            return f"{d.month}.{d.day}"
        except (ValueError, TypeError):
            continue
    return at or ""


COLUMNS = [
    ("序号",              lambda o, i: i,                 6),
    ("时间",              lambda o, i: _day(o.at),        9),
    ("新增/存量/单C",      lambda o, i: o.kind,           14),
    ("套餐资费",           lambda o, i: o.fee,            14),
    # 客户手机号，机器人一律脱敏不落盘。列保留是为了和原表对齐，方便人工补
    ("设备号",            lambda o, i: None,             14),
    ("新增宽带",           lambda o, i: _num(o.broadband), 10),
    ("新增中高套",         lambda o, i: _num(o.high),      11),
    ("长续约",            lambda o, i: _num(o.renew),      9),
    ("千兆/500兆",        lambda o, i: _num(o.giga),      11),
    ("单号卡",            lambda o, i: _num(o.card),       9),
    # FTTR 不单列（09-29 用户：FTTR 在攻坚群总表里算套包，活动群不统计）
    ("云盒/AI视联网/天翼智屏",  lambda o, i: _num(o.box),  20),
    ("价值积分",           lambda o, i: o.points,         10),
    ("揽收人",            lambda o, i: o.collector,       11),
    ("战区",              lambda o, i: o.team,           12),
    ("小组",              lambda o, i: o.group,          18),
    ("备注",              lambda o, i: o.note,           16),
]


def write_orders_sheet(ws, orders: list[Order]) -> None:
    """把订单写进一个工作表。样式跟着原表的观感走：橙色表头 + 细边框。"""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="EE602C")
    head_font = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="D9DDE3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # 待核对的行标黄，扫一眼就知道哪几行要人工补
    warn_fill = PatternFill("solid", fgColor="FFF3CD")

    for c, (name, _fn, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(1, c, name)
        cell.fill, cell.font, cell.alignment, cell.border = (
            head_fill, head_font, center, border)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 28

    for i, o in enumerate(orders, 1):
        for c, (_name, fn, _w) in enumerate(COLUMNS, 1):
            cell = ws.cell(i + 1, c, fn(o, i))
            cell.alignment = center
            cell.border = border
            if o.needs_review:
                cell.fill = warn_fill

    # 冻结表头，几十行往下翻时还看得见列名
    ws.freeze_panes = "A2"


# ---------- 晒单数据 ----------

# 左块：每个小组一行，和原表「晒单数据」A:L 一致
SUMMARY_COLUMNS = [
    ("战区", 11), ("小组", 7), ("成员", 18),
    ("新增宽带", 10), ("中高套", 9), ("长续约", 9), ("千兆/500M", 11),
    ("单号卡", 9), ("云盒/AI视联网/天翼智屏", 22), ("价值积分（分）", 14),
    ("综合完成进度", 14), ("排名", 7),
]
# 右块：目标值 + 战队进度，和原表 O:X 一致
TARGET_COLUMNS = [
    ("战区", 11), ("新增宽带（25%）", 15), ("中高套（30%）", 14), ("长续约", 9),
    ("千兆/500兆", 11), ("单号卡", 9), ("云盒/AI视联网/天翼智屏（15%）", 28),
    ("价值积分（20%）", 15), ("综合进度", 11), ("排名", 7),
]
# 右块从这一列开始（留两列空白隔开）
_TARGET_START_COL = len(SUMMARY_COLUMNS) + 3


def _rank_cell(ts: TeamStat | None):
    """名次格：没参与排名留空，0 分写「-」，有分才写数字。"""
    if ts is None or not ts.rank:
        return ""
    return ts.rank if ts.progress > 0 else "-"


def summary_title(period: str, title: str, until: datetime | None = None) -> str:
    """「8.27-8.30日各支局晒单情况（截至8月28日20：55）」。"""
    until = until or datetime.now()
    return f"{period}日各支局晒单情况（截至{until.month}月{until.day}日{until:%H:%M}）"


def write_summary_sheet(ws, groups: list[GroupStat], teams: list[TeamStat],
                        targets: Targets, title: str) -> None:
    """
    「晒单数据」表：小组明细 + 战队合计 + 综合进度 + 排名，加右侧目标值块。

    全部写结果值不写公式（理由见模块开头）。
    战区 / 综合完成进度 / 排名 三列在每个战队块内纵向合并，和原表观感一致。
    """
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="EE602C")
    head_font = Font(bold=True, color="FFFFFF", size=11)
    total_fill = PatternFill("solid", fgColor="FDE9D9")
    total_font = Font(bold=True)
    title_font = Font(bold=True, size=14)
    thin = Side(style="thin", color="D9DDE3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    n_left = len(SUMMARY_COLUMNS)

    # 标题行：横跨左块
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_left)
    c = ws.cell(1, 1, title)
    c.font, c.alignment = title_font, Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    # 左块表头（第 2 行）
    for i, (name, width) in enumerate(SUMMARY_COLUMNS, 1):
        cell = ws.cell(2, i, name)
        cell.fill, cell.font, cell.alignment, cell.border = head_fill, head_font, center, border
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[2].height = 30

    by_team = {t.team: t for t in teams}
    row = 3
    current = None
    block_start = row
    ranked_teams = [t.team for t in teams if t.rank]

    def close_block(team: str, start: int, end: int) -> int:
        """写战队合计行，并把战区 / 进度 / 排名三列在块内合并。返回下一行号。"""
        ts = by_team.get(team)
        vals = ["", "合计", ""]
        if ts is not None:
            vals += [ts.broadband, ts.high, ts.renew, ts.giga, ts.card, ts.box, ts.points]
        else:
            vals += [0] * 7
        vals += ["", ""]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(end + 1, i, v)
            cell.fill, cell.font, cell.alignment, cell.border = total_fill, total_font, center, border
        last = end + 1
        # 合并：战区列、综合完成进度列、排名列
        for col, value, fmt in (
            (1, team, None),
            # 「未归队」不参与排名，进度和名次都留空；0 分的战队名次写「-」
            (n_left - 1, ts.progress if (ts and ts.rank) else "", "0.00%"),
            (n_left, _rank_cell(ts), None),
        ):
            ws.merge_cells(start_row=start, start_column=col, end_row=last, end_column=col)
            cell = ws.cell(start, col, value)
            cell.alignment, cell.border = center, border
            if fmt:
                cell.number_format = fmt
            if col == 1:
                cell.font = Font(bold=True)
        return last + 1

    for g in groups:
        if g.team != current:
            if current is not None:
                row = close_block(current, block_start, row - 1)
            current, block_start = g.team, row
        vals = ["", g.label, g.members, g.broadband, g.high, g.renew, g.giga,
                g.card, g.box, g.points, "", ""]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row, i, v)
            cell.alignment, cell.border = center, border
        row += 1
    if current is not None:
        row = close_block(current, block_start, row - 1)

    # 右块：目标值 + 进度 + 排名
    c0 = _TARGET_START_COL
    for i, (name, width) in enumerate(TARGET_COLUMNS):
        cell = ws.cell(1, c0 + i, name)
        cell.fill, cell.font, cell.alignment, cell.border = head_fill, head_font, center, border
        ws.column_dimensions[get_column_letter(c0 + i)].width = width
    for r, team in enumerate(ranked_teams, 2):
        ts = by_team[team]
        vals = [team, targets.broadband, targets.high, None, None, None,
                targets.box, targets.points, ts.progress, _rank_cell(ts)]
        for i, v in enumerate(vals):
            cell = ws.cell(r, c0 + i, v)
            cell.alignment, cell.border = center, border
            if i == 8:
                cell.number_format = "0.00%"

    ws.freeze_panes = "A3"


def build_workbook(camp: Campaign, out: str | Path = "业绩清单.xlsx", title: str = "",
                   rank: tuple[str, str] = ("", ""),
                   ledger: tuple[str, str] | None = None,
                   checkin_sheets: list | None = None, checkin_columns: int = 25) -> Path:
    """
    生成 Excel：「晒单数据」（小组汇总 + 排名）+「清单」（每单一行）。

    两张表的时间范围【可以不同】：排名按天比、清单按周记，所以分开传。
    rank / ledger 都是 (since, until) 的完整时刻，能切出「上午」「下午」这种时段 ——
    表里显示的时间列只到天，但筛选用的是完整时刻。ledger 不传就和 rank 一样。
    """
    from openpyxl import Workbook

    r_since, r_until = rank
    l_since, l_until = ledger if ledger is not None else rank
    orders = [o for o in camp.orders
              if (not l_since or o.at >= l_since) and (not l_until or o.at <= l_until)]

    wb = Workbook()
    ws = wb.active
    ws.title = "晒单数据"
    write_summary_sheet(ws, camp.group_stats(r_since, r_until), camp.stats(r_since, r_until),
                        camp.targets, title or summary_title("", "各支局"))
    ws2 = wb.create_sheet("清单")
    write_orders_sheet(ws2, orders)
    # 打卡表：按时段各一张（[(sheet 名, 标题, 记录)]，由 checkin.day_sheets 准备）
    if checkin_sheets:
        from .checkin import write_checkin_sheet
        for name, sheet_title, records in checkin_sheets:
            write_checkin_sheet(wb.create_sheet(name), records, camp.roster, sheet_title, checkin_columns)

    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    log.info("已生成 %s：清单 %d 单（%s ~ %s），排名按 %s ~ %s",
             out.name, len(orders), l_since[:10] or "开始", l_until[:10] or "现在",
             r_since[:10] or "开始", r_until[:10] or "现在")
    return out
