"""
支局长打卡点检表（09-29 国庆版，按用户给的「9.28日支局长打卡情况点检」图），只发攻坚群。
规则见 docs/国庆多群版功能说明.md 第八节（用户 09-29 确认）：

- 07:30 起，每一张单聊打卡照（不管谁发的、这张打卡记没记）都和支局长的人脸比对；
  认出谁就记「谁、几点（水印时间）、照片里几张脸、工作内容」。认不出支局长的照片什么都不存
- 早会：07:30–09:00 里第一张认出他的【集合照】；中午集合：13:30–14:30 里第一张集合照
  集合照 = 照片里 ≥3 张脸，或者水印工作内容写了 早会 / 集合 / 晨会 / 例会
  窗口里认出了他但不是集合照 → 「自行打卡」（红）；窗口过了一张都没有 → 「未打卡」（红）
- 上门：其余认出他的照片按时间填「第1户、第2户…」；同一个人 10 分钟内的几张算同一户
- 休息：管理员单聊「休息 杨钰航」（可带日期 10.2）→ 早会、中午集合、第1户写「休息」（黄）
- 补记：管理员单聊「补记 严健齐 10:37」（可带支局、日期）→ 加一户
- 名字标红：14:00 以后一户都没有（当天要到 18:00 以后才判，之前的日子直接判；休息的不标）

只存认出来的支局长和时间；人脸特征在 faces.py（只存特征不存照片）。
"""

from __future__ import annotations

import io
import json
import logging
import re
import threading
from dataclasses import dataclass, field
from datetime import date, datetime, time as dtime
from pathlib import Path

from .paths import DATA

log = logging.getLogger("botcore.chiefs")

STORE_FILE = "chief_sightings.json"
START = "07:30"
EARLY = ("07:30", "09:00")
NOON = ("13:30", "14:30")
AFTERNOON = "14:00"          # 这之后一户都没有 → 名字标红
RED_JUDGE_AT = "18:00"        # 当天这个点之后才判标红
MERGE_MINUTES = 10
MIN_COLUMNS = 9
_GATHER_WORDS = re.compile(r"早会|集合|晨会|例会")

RED = (255, 0, 0)
YELLOW = (255, 255, 0)


def _hm(t: str) -> str:
    """08:20 -> 8:20（和样表一样不补 0）。"""
    h, m = t.split(":")
    return f"{int(h)}:{m}"


def _minutes(t: str) -> int:
    h, m = t.split(":")
    return int(h) * 60 + int(m)


def is_gathering(faces: int, work: str) -> bool:
    return faces >= 3 or bool(_GATHER_WORDS.search(work or ""))


@dataclass
class ChiefRow:
    cls: str
    branch: str
    name: str
    early: str = ""           # 时间 / 自行打卡 / 未打卡 / 休息 / ""
    noon: str = ""
    visits: list = field(default_factory=list)      # ["10:58", ...] 或 ["休息"]
    red_name: bool = False
    rest: bool = False


class ChiefBoard:
    def __init__(self, faces, rosters, store, data_dir: Path = DATA) -> None:
        self.faces = faces            # faces.FaceIndex
        self.rosters = rosters        # regions.RosterSet（支局长名单在 rosters.chiefs）
        self.store = store
        self.path = Path(data_dir) / STORE_FILE
        self._lock = threading.Lock()
        self._data = self._load()
        # 人脸比对自己一条队列（09-29 用户提议）：打卡那边读完水印就把照片交过来，不等比对完就继续判打卡类型、
        # 发回执；几个人同时发也是一张张排队比，不会漏。单线程：OpenCV 人脸模型不保证线程安全，2 核机器也够
        from concurrent.futures import ThreadPoolExecutor
        self._pool = ThreadPoolExecutor(max_workers=1, thread_name_prefix="chief-face")

    # ---------- 存储 ----------

    def _load(self) -> dict:
        try:
            raw = json.loads(self.path.read_text(encoding="utf-8"))
            return {"sightings": list(raw.get("sightings") or []), "manual": list(raw.get("manual") or [])}
        except FileNotFoundError:
            return {"sightings": [], "manual": []}
        except Exception as exc:
            log.error("支局长点检记录读不了，当作空的：%s", exc)
            return {"sightings": [], "manual": []}

    def _save(self) -> None:
        tmp = self.path.with_suffix(".tmp")
        tmp.write_text(json.dumps(self._data, ensure_ascii=False, indent=1), encoding="utf-8")
        tmp.replace(self.path)

    # ---------- 每张打卡照 ----------

    def submit(self, blob: bytes, wm: dict, photo: str):
        """打卡那边调这个：交给人脸队列，马上返回（打卡判定和人脸比对同时跑）。"""
        def job():
            try:
                self.on_photo(blob, wm, photo)
            except Exception:
                log.exception("支局长人脸点检出错（%s）", wm.get("stamp"))
        try:
            return self._pool.submit(job)
        except RuntimeError:
            return None

    def wait_idle(self, timeout: float = 60) -> None:
        """等队列里的比对都做完（出点检表前、测试用）。"""
        try:
            self._pool.submit(lambda: None).result(timeout=timeout)
        except Exception:
            pass

    def on_photo(self, blob: bytes, wm: dict, photo: str) -> list[str]:
        """CheckinRecorder 每张有水印时间的照片都调。返回认出的支局长（测试用）。"""
        from . import regions
        if not regions.multi_enabled(self.store):
            return []
        self.rosters.chiefs.maybe_reload()
        chiefs = self.rosters.chiefs.names()
        stamp = wm.get("stamp") or ""
        if not chiefs or len(stamp) < 16 or stamp[11:16] < START or not self.faces.available():
            return []
        r = self.faces.identify(blob)
        found = [n for n, _s in r["matches"] if n in chiefs]
        if not found:
            return []
        day, hm = stamp[:10], stamp[11:16]
        with self._lock:
            have = {(s["photo"], s["name"]) for s in self._data["sightings"]}
            for name in found:
                if (photo, name) in have:
                    continue
                self._data["sightings"].append({"day": day, "time": hm, "name": name, "faces": r["faces"],
                                                "work": wm.get("work", ""), "photo": photo})
            self._save()
        log.info("支局长点检 | %s %s | 认出 %s（照片里 %d 张脸）", day, hm, "、".join(found), r["faces"])
        return found

    def add_manual(self, kind: str, name: str, day: str, hm: str = "") -> None:
        with self._lock:
            self._data["manual"].append({"kind": kind, "name": name, "day": day, "time": hm,
                                         "at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")})
            self._save()

    # ---------- 出表 ----------

    def rows(self, day: str, now: datetime) -> list[ChiefRow]:
        chiefs = self.rosters.chiefs.chiefs
        today = now.strftime("%Y-%m-%d")
        past = day < today
        cur = now.strftime("%H:%M") if not past else "23:59"
        with self._lock:
            sights = [s for s in self._data["sightings"] if s["day"] == day]
            manual = [m for m in self._data["manual"] if m["day"] == day]
        out = []
        for c in chiefs:
            name = c["name"]
            row = ChiefRow(c.get("class", ""), c.get("branch", ""), name)
            mine = sorted((s for s in sights if s["name"] == name), key=lambda s: s["time"])
            rest = any(m["kind"] == "rest" and m["name"] == name for m in manual)
            if rest:
                row.rest = True
                row.early = row.noon = "休息"
                row.visits = ["休息"]
                out.append(row)
                continue
            used = set()

            def gather(win):
                lo, hi = win
                inside = [s for s in mine if lo <= s["time"] <= hi]
                gs = [s for s in inside if is_gathering(s.get("faces", 0), s.get("work", ""))]
                for s in gs:
                    used.add(id(s))                 # 集合照不算上门
                if gs:
                    return _hm(gs[0]["time"])
                if inside:
                    return "自行打卡"
                return "未打卡" if cur > hi else ""

            row.early = gather(EARLY)
            row.noon = gather(NOON)
            times = sorted([s["time"] for s in mine if id(s) not in used] +
                           [m["time"] for m in manual if m["kind"] == "add" and m["name"] == name and m["time"]])
            visits = []
            for t in times:
                if visits and _minutes(t) - _minutes(visits[-1]) < MERGE_MINUTES:
                    continue
                visits.append(t)
            row.visits = [_hm(t) for t in visits]
            judge = past or cur >= RED_JUDGE_AT
            row.red_name = judge and not any(t >= AFTERNOON for t in visits)
            out.append(row)
        return out

    def title(self, day: str, now: datetime) -> str:
        d = datetime.strptime(day, "%Y-%m-%d")
        end = now.strftime("%H:%M") if day == now.strftime("%Y-%m-%d") else "21:00"
        return f"{d.month}.{d.day}日支局长打卡情况点检（{_hm(START)}-{_hm(end)}）"

    def render_png(self, day: str, now: datetime) -> bytes:
        from PIL import Image, ImageDraw
        from .bulletin import _font
        rows = self.rows(day, now)
        ncol = max(MIN_COLUMNS, max((len(r.visits) for r in rows), default=0))
        f_title, f_head, f_cell = _font(26, True), _font(18, True), _font(18)
        widths = [80, 100, 110, 150, 150] + [110] * ncol
        width = sum(widths)
        h_title, h_head, h_row = 48, 56, 36
        height = h_title + h_head + h_row * max(len(rows), 1) + 2
        img = Image.new("RGB", (width, height), "white")
        d = ImageDraw.Draw(img)
        line = (60, 60, 60)

        def box(x0, y0, x1, y1, text="", font=f_cell, fill=None):
            d.rectangle([x0, y0, x1, y1], fill=fill, outline=line)
            if text:
                tw = d.textlength(text, font=font)
                d.text((x0 + (x1 - x0 - tw) / 2, y0 + (y1 - y0 - font.size) / 2 - 1), text, font=font, fill=(0, 0, 0))

        box(0, 0, width - 1, h_title, self.title(day, now), f_title)
        heads = ["类别", "支局", "支局长", "早会", "中午集合"] + [f"第{i}户" for i in range(1, ncol + 1)]
        x = 0
        for h, w in zip(heads, widths):
            box(x, h_title, x + w, h_title + h_head, h, f_head)
            x += w
        y0 = h_title + h_head

        def spans(key):
            i = 0
            while i < len(rows):
                j = i
                while j < len(rows) and key(rows[j]) == key(rows[i]):
                    j += 1
                yield i, j
                i = j

        for i, j in spans(lambda r: r.cls):
            box(0, y0 + i * h_row, widths[0], y0 + j * h_row, rows[i].cls, f_head)
        for i, j in spans(lambda r: (r.cls, r.branch)):
            box(widths[0], y0 + i * h_row, widths[0] + widths[1], y0 + j * h_row, rows[i].branch, f_head)
        for r_i, r in enumerate(rows):
            y = y0 + r_i * h_row
            x = widths[0] + widths[1]
            cells = [(r.name, RED if r.red_name else None)]
            for v in (r.early, r.noon):
                cells.append((v, YELLOW if v == "休息" else RED if v in ("未打卡", "自行打卡") else None))
            for k in range(ncol):
                v = r.visits[k] if k < len(r.visits) else ""
                cells.append((v, YELLOW if v == "休息" else None))
            for (v, fill), w in zip(cells, widths[2:]):
                box(x, y, x + w, y + h_row, v, f_cell, fill)
                x += w
        buf = io.BytesIO()
        img.save(buf, "PNG")
        return buf.getvalue()

    def write_sheet(self, ws, day: str, now: datetime) -> None:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        rows = self.rows(day, now)
        ncol = max(MIN_COLUMNS, max((len(r.visits) for r in rows), default=0))
        heads = ["类别", "支局", "支局长", "早会", "中午集合"] + [f"第{i}户" for i in range(1, ncol + 1)]
        ws.cell(1, 1, self.title(day, now))
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(heads))
        for c, h in enumerate(heads, 1):
            ws.cell(2, c, h)
        red, yellow = PatternFill("solid", fgColor="FF0000"), PatternFill("solid", fgColor="FFFF00")
        for i, r in enumerate(rows):
            vals = [r.cls, r.branch, r.name, r.early, r.noon] + [r.visits[k] if k < len(r.visits) else ""
                                                                  for k in range(ncol)]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(3 + i, c, v)
                if (c == 3 and r.red_name) or v in ("未打卡", "自行打卡"):
                    cell.fill = red
                elif v == "休息":
                    cell.fill = yellow
        thin = Side(style="thin", color="666666")
        for row in ws.iter_rows(min_row=1, max_row=2 + len(rows), max_col=len(heads)):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.cell(1, 1).font = Font(bold=True, size=14)
        for col_start, key in ((1, lambda r: r.cls), (2, lambda r: (r.cls, r.branch))):
            i = 0
            while i < len(rows):
                j = i
                while j < len(rows) and key(rows[j]) == key(rows[i]):
                    j += 1
                if j - i > 1:
                    ws.merge_cells(start_row=3 + i, start_column=col_start, end_row=2 + j, end_column=col_start)
                i = j

    def send(self, platform, chat: str, now: datetime) -> bool:
        if not self.rosters.chiefs.chiefs:
            log.info("还没导支局长名单，点检表不发")
            return True
        self.wait_idle(30)            # 刚到的几张还在比对：等一下再出表，别漏掉
        day = now.strftime("%Y-%m-%d")
        png = self.render_png(day, now)
        out = DATA / self.store.bulletin_settings().out_dir
        out.mkdir(parents=True, exist_ok=True)
        name = f"支局长点检_{now:%m%d_%H%M}.png"
        (out / name).write_bytes(png)
        return bool(platform.send_image(chat, png, name)) if hasattr(platform, "send_image") else False


def build_excel(store, camp, day: date, out: Path) -> Path:
    """存档 / 单聊「支局长表格」：那天的点检表。"""
    from openpyxl import Workbook
    from . import faces, regions
    board = ChiefBoard(faces.FaceIndex(), regions.RosterSet(), store)
    now = datetime.now()
    at = now if day >= now.date() else datetime.combine(day, dtime(23, 59))
    wb = Workbook()
    wb.active.title = "支局长点检"
    board.write_sheet(wb.active, day.strftime("%Y-%m-%d"), at)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
