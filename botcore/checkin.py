"""
打卡核验：每张「今日水印」照片算一次打卡，判出类型，按时段出「打卡」表。

判定顺序（2026-09-24 和业务确认的口径）：
    ① 水印「工作内容」里有办成的业务（新增 / 续签 / 改档 / 号卡 …）
       → 不看照片，直接记成这笔业务，如「129改149」「新增宽带」「19山城卡」   填绿
       「拆机维系」后面写了成功的也算业务                                        填绿
    ② 工作内容是「异网打卡」「拆机维系」这类工作类型（config: keep_as_is）
       → 照原样写进格子，不看照片                                                不填色
    ③ 其余交给视觉模型看照片：
         拦截      有人在交谈（户外、街边、摊位、营业厅、店里）；摊位上摆着资料设备也算拦截  不填色
         设备演示  电视上是云盒 / IPTV 菜单界面，或拿着设备给客户讲；装机、普通节目不算        不填色
         板夹      蓝色板夹 / 资费单拿给客户看；桌上随手放着的不算                            不填色
         算账      手写算式、计算器、手机业务办理界面、拍证件办业务              不填色
         异网      手机屏幕上是移动 / 联通 / 广电等其他运营商的套餐资费界面      不填色
         无效      看不到人；或入户 / 办公室里只是有人在场、没做上面的展示动作   填红
                   （09-24 用户定：上门入户、客户在场但没展示东西 = 无效，不是拦截）
       一张照片符合多类时任选一类（业务定的）
       模型的结果再用本机 OCR 核一遍（refine）：
         判「异网」要在画面里读到移动 / 联通 / 广电等字样，读不到改算账
         判「算账」要在画面里读到办理类的字或数字算式，读不到退回拦截
         （实测模型会看着营业厅背景脑补「电信业务界面」、把座机当计算器）
    ④ 视觉模型没配 / 调不通 → 记「待判定」，不填色，不耽误出表

时段（按水印时间，不按发送时间）：上午 08:30-11:50 / 下午 14:00-18:00 / 晚上 18:00-21:00。
    时段外算下一个时段（08:30 前 → 上午，11:50-14:00 → 下午）；21:00 以后不统计（记录留着，表和统计里不出现）
表格式照原表「城区支局活动8.27-8.30.xlsx」的「打卡」sheet：
    行 = 战区 → 小组 → 成员；列 = 第1户 … 第25户，按时间往右排
    格子「18:57（129改149）」；成员名标红 = 这个小组一条打卡都没有；合计、触点总量含无效

隐私：
    · 送视觉模型前，先在本机把照片里的证件、验证码、病历、长串数字那几行涂黑，再缩到 1024 以内
    · 照片本身不落盘，只存判定结果；同一张照片按水印防伪码去重（企微转发会重压缩，文件指纹对不上）
"""

from __future__ import annotations

import base64
import hashlib
import io
import json
import logging
import os
import re
import threading
import time
from dataclasses import asdict, dataclass, field
from datetime import datetime, timedelta
from pathlib import Path

import requests

from .distill import _CHROME, _MEMBERS, extract_branch, extract_stamp, redact
from .image import _ENGINE, OCRSettings
from .backup import keep_version
from .paths import DATA

log = logging.getLogger("botcore.checkin")

GREEN, RED = "green", "red"

# 视觉模型的类别 -> 格子里写的字
# 设备演示在通报表里写「云盒打卡」（09-24 用户给的通报样表就是这么写的）
DISPLAY = {"拦截": "现场拦截", "设备演示": "云盒打卡", "板夹": "板夹打卡", "算账": "算账打卡",
           "异网": "异网打卡", "无效": "无效打卡"}


@dataclass
class CheckinSettings:
    enabled: bool = False
    slots: list = field(default_factory=lambda: [
        {"name": "上午", "start": "08:30", "end": "11:50"},
        {"name": "下午", "start": "14:00", "end": "18:00"},
        {"name": "晚上", "start": "18:00", "end": "21:00"},
    ])
    # 工作内容里写了这些，照原样写进格子、不看照片。
    # 可以写成 {match: 搬家打卡, label: 板夹打卡}：工作内容里有「搬家打卡」（「板夹」的同音错字），格子里写「板夹打卡」
    keep_as_is: list = field(default_factory=lambda: ["异网打卡", "拆机维系",
                                                      {"match": "搬家打卡", "label": "板夹打卡"}])
    columns: int = 13                 # 表里至少留几户的列（通报样表是第 1 ~ 13 户）
    # 表格标题里的活动日期，如「9.24-9.27」；留空就写当天「9.24」
    span: str = ""
    # ---- 打卡通报（09-24 用户定的时刻），发「文字 + 打卡表图片」----
    bulletin_enabled: bool = False
    bulletin_times: list = field(default_factory=lambda: [
        "09:30", "10:30", "11:50", "15:00", "16:00", "17:00", "18:00", "19:30", "20:50"])
    bulletin_days: list = field(default_factory=lambda: [1, 2, 3, 4, 5, 6, 7])
    # 群里 @机器人 说这些，立刻发一份当前时段的打卡通报
    bulletin_commands: list = field(default_factory=lambda: ["打卡通报", "打卡情况"])
    bulletin_state_file: str = "checkin_bulletin_state.json"
    store_file: str = "checkin.json"
    image_max_side: int = 1024        # 送模型前缩到多大
    vision_base_url: str = "https://dashscope.aliyuncs.com/compatible-mode/v1"
    vision_model: str = "qwen-vl-max"
    vision_api_key_env: str = "VISION_API_KEY"
    vision_timeout: float = 30.0
    vision_max_calls_per_minute: int = 20
    # 当前模型免费额度用完 / 没开通时按顺序换下一个（09-28 用户定：不用测，额度用完直接换）
    vision_fallback_models: list = field(default_factory=lambda: ["qwen3-vl-flash", "qwen-vl-plus"])
    pending_keep_days: int = 3        # 「待判定」的照片（涂黑缩小后的）最多留几天等补判
    retry_minutes: int = 10           # 每隔几分钟补判一次「待判定」


@dataclass
class Checkin:
    at: str                 # 水印时间 YYYY-MM-DD HH:MM
    slot: str               # 上午 / 下午 / 晚上 / ""（时段外）
    branch: str             # 水印里的支局
    members: str            # 水印里的小组成员（原样）
    work: str               # 水印里的工作内容（已脱敏）
    label: str              # 格子括号里的字
    color: str              # green / red / ""
    source: str             # 业务 / 原样 / 视觉 / 待判定
    evidence: str = ""      # 视觉模型说它看到了什么
    photo: str = ""         # 照片内容指纹，去重用
    msg_id: str = ""
    received: str = ""
    sender: str = ""        # 发照片的人的企微账号 ID（长连接只给 ID 不给姓名）
    code: str = ""          # 水印防伪码：每张照片唯一，去重用（企微转发会重压缩，文件指纹对不上）
    model: str = ""         # 哪个视觉模型判的（额度用完会自动换模型，留个底）
    region: str = ""        # 多群：城区 / 农村（按发送人 / 水印小组成员认出来的活动人员所在区；老记录为空 = 城区）


# ---------------- 水印 ----------------

_PLACEHOLDER = re.compile(r"请输入内容|请输入|点击输入")
# 不能用 distill 的 _WORK_FIELD：它为了晒单提纯，遇到「打卡」就截断，
# 「异网打卡」会变成「异网」、「上门打卡」变成「上门」，原样照写就对不上了
_STOP = r"(?=\s*(?:小组成员|拍摄人|工作单位|时\s*间|地\s*点|经纬度|海拔|支\s*局|防伪|备\s*注|验真|$))"
_WORK = re.compile(r"工作内容\s*[:：]?\s*(.{0,80}?)" + _STOP)
_REMARK = re.compile(r"备\s*注\s*[:：]?\s*(.{0,80}?)" + _STOP)


_CODE = re.compile(r"防\s*伪\s*[:：]?\s*([A-Za-z0-9]{8,24})")


def is_watermark(text: str) -> bool:
    """今日水印相机的照片：有「今日水印」字样，或者同时有「工作内容」和成员字段。"""
    if "今日水印" in text or "真实可验" in text:
        return True
    return "工作内容" in text and ("小组成员" in text or "拍摄人" in text)


def parse_watermark(text: str) -> dict:
    clean = _CHROME.sub(" ", redact(text))
    m = _WORK.search(clean) or _REMARK.search(clean)
    work = m.group(1).strip(" ：:，,。.") if m else ""
    mem = _MEMBERS.search(clean)
    members = mem.group(1).strip() if mem else ""
    if _PLACEHOLDER.search(members):
        members = ""
    if _PLACEHOLDER.search(work) or work in ("相机", "真实可验"):
        work = ""
    code = _CODE.search(text)
    return {"stamp": extract_stamp(clean), "branch": extract_branch(text),
            "members": members, "work": work, "code": code.group(1).upper() if code else ""}


def slot_of(hhmm: str, slots: list) -> str:
    """
    按水印时刻落时段（09-24 业务定的）：
      · 在某个时段里的就算那个时段；时段之间的空档（11:50-14:00）和第一个时段之前，算【下一个】时段
      · 最后一个时段结束之后（21:00 以后）不统计，返回空串；21:00 整还算晚上
    """
    ordered = sorted(slots or (), key=lambda x: str(x.get("start", "")))
    for i, s in enumerate(ordered):
        end = str(s.get("end", ""))
        nxt = str(ordered[i + 1].get("start", "")) if i + 1 < len(ordered) else ""
        # 结束那一分钟算本时段（11:50 算上午、21:00 算晚上）；和下一段首尾相接的（18:00）算下一段
        if hhmm < end or (hhmm == end and nxt != end):
            return str(s.get("name", ""))
    return ""


# ---------------- ① 业务优先 ----------------

_NEG = re.compile(r"预约|意向|咨询|了解|回访|考虑|未办|没办|宣传")
_SEP = r"\s，,。；;、"
_BIZ = [
    # 档位变化：129改149 / 139升199+10 / 109提值169
    re.compile(r"\d{2,4}\s*(?:改|升档|升级|升|提值|→|->)\s*\d{2,4}(?:\s*\+\s*\d{1,3})*"),
    # 续签：199续签 / 139续约 / 99+10云盒续签（「续签」前面可以隔几个字）
    re.compile(r"\d{2,4}(?:\s*\+\s*\d{1,3})*[^\d\s，,。；;、]{0,4}?(?:续签|续约|续费)"),
    # 号卡类：19山城卡 / 优质号卡 / 橙分期
    re.compile(rf"[^{_SEP}]{{0,4}}(?:山城卡|号卡|橙分期|携转|单C|副卡|副宽)"),
    # 新增宽带 / 新装 / 新办
    re.compile(rf"(?:新增|新装|新办)[^{_SEP}]{{0,10}}"),
    # 办理 xx / 开通 xx / 加装 xx
    re.compile(rf"(?:办理|开通|加装|受理|装机)[^{_SEP}]{{1,10}}"),
    # 策反异网：算新增（09-24 用户定）。「异网策反139套餐」-> 「新增139」
    re.compile(rf"策反[^{_SEP}]{{0,10}}"),
]
_LEAD = re.compile(r"^(?:号码回归|回归|办理|开通|加装|受理|装机|新办理|理|码)+")
_TAIL = re.compile(r"(?:拍单|成功|办理|一户|一单|了|的)+$")
_MAINTAIN_OK = re.compile(r"拆机维系.{0,6}(?:成功|挽留|已维系|留住)|(?:成功|挽留).{0,4}拆机维系")


def business_of(work: str) -> str:
    """工作内容里办成的业务，返回写进格子的短语；没有返回空串。"""
    if not work:
        return ""
    # OCR 常在字中间插空格（「融合宽 带」），不去掉会截成「新增199融合宽」
    work = re.sub(r"(?<=[\u4e00-\u9fa5\w])\s+(?=[\u4e00-\u9fa5\w])", "", work)
    for pat in _BIZ:
        for m in pat.finditer(work):
            # 「预约新增」「意向办理」这种没办成：业务词前面紧挨着否定词的不算
            if _NEG.search(work[max(0, m.start() - 3):m.start()]):
                continue
            if m.group(0).startswith("策反"):
                tier = re.search(r"\d{2,4}", m.group(0))
                return "新增" + (tier.group(0) if tier else "")
            phrase = _TAIL.sub("", _LEAD.sub("", re.sub(r"\s+", "", m.group(0))))
            if len(phrase) >= 2:
                return phrase[:12]
    return ""


def rule_label(work: str, keep_as_is: list) -> tuple[str, str, str] | None:
    """①② 两步：(label, color, source)；要看照片返回 None。"""
    if _MAINTAIN_OK.search(work or ""):
        return "拆机维系", GREEN, "业务"
    biz = business_of(work)
    if biz:
        return biz, GREEN, "业务"
    for item in keep_as_is or ():
        word, label = ((item.get("match", ""), item.get("label", "")) if isinstance(item, dict)
                       else (str(item), str(item)))
        if word and word in (work or ""):
            return label or word, "", "原样"
    return None


# ---------------- ③ 视觉模型 ----------------

PROMPT = """这是中国电信一线营销人员用「今日水印相机」拍的上门/外出打卡照片。照片上叠加的水印文字不作为依据，只看画面。
按下面的步骤判断：
第 1 步：是不是对着手机屏幕拍的特写？（屏幕占画面一大半、屏幕上的字和界面看得清；这一步不要求画面里有人）
    手机只是拿在手里、拍得小或拍糊了、屏幕内容看不清的，不算，直接跳到第 2 步。
    · 屏幕上看得到中国移动、中国联通、中国广电等【电信以外】运营商的名字或标志，显示套餐、资费、话费账单、
      流量余额或营业厅 App 界面 → 异网
    · 屏幕上是中国电信的业务办理、受理、下单、套餐确认界面 → 算账
    · 都不是就往下看
第 2 步：画面里有没有人？哪怕只拍到腿、手、背影、半个身子也算有人。
    没有任何人（空走廊、空房间、门牌、风景）→ 无效
第 3 步：有人的话：
    · 在摊位上（帐篷、折叠桌、展台、易拉宝）→ 拦截（摊位桌上摆着宣传单、板夹、设备都还是拦截）
    · 电视屏幕上是云盒 / IPTV 的【菜单界面】（一格格的栏目、海报、应用图标），或者有人把光猫、云盒、路由器拿在手里给客户看、讲解 → 设备演示
      （电视在播普通节目（新闻、央视台标、体育比赛、电视剧画面）不算；监控摄像头的多格画面不是云盒菜单；
       工作人员在弱电箱、墙上、电视柜后面装机接线调试不算）
    · 有人把【蓝色塑料板夹】或资费单、宣传单拿在手里 / 翻开递到客户面前看 → 板夹（桌上压着、随手放着、没人给客户看的单子不算）
    · 纸上手写的算式或金额、工作人员拿笔在纸上（宣传单、资费单、白纸）给客户写写画画圈圈算算、屏幕亮着数字的计算器、
      手机上的【电信】业务办理界面、正在拍证件 → 算账（座机电话、遥控器都不是计算器）
    · 上面都没有：
        在户外、街边、小区、营业厅、店里，有人在和客户 / 路人交谈 → 拦截
        没人在交谈（工作人员自己看手机、自己干活、只拍到手臂桌面）→ 无效
        在客户家里（入户）或办公室里，只是有人在场、在聊天，或者工作人员在装机接线 → 无效
规则：只根据画面里清清楚楚看到的东西判断；看不清的不要猜。一张照片符合多类时选最明显的一类。
照片里涂黑的方块是隐私遮挡，忽略即可。
只输出一个 JSON：{"people": true或false, "place": "摊位/户外/营业厅/店铺/入户/办公室/其他 之一", "tv": "云盒菜单/普通节目/没有电视 之一", "installing": 工作人员在装机接线调试时 true 否则 false, "folder": 画面里有蓝色板夹或资费单翻开摆在 / 递到客户面前时 true 否则 false, "talking": 画面里有人在和客户或路人交谈、互动时 true（工作人员自己看手机、干活、只拍到手臂都算 false）, "label": "异网/无效/设备演示/板夹/算账/拦截 之一", "evidence": "画面里看到了什么，20 字以内"}"""

# 异网要有文字证据：模型认得出「这是手机屏幕」，但分不清哪家运营商（实测把电信营业厅界面也判成异网）。
# 本机 OCR 在屏幕上读到其他运营商的字样才算异网，否则按「手机办理界面」改判算账。
# 水印里只有「中国电信」，不会误中
_OTHER_CARRIER = re.compile(r"移动|联通|广电|全球通|动感地带|神州行|CMCC|China\s*Mobile|Unicom", re.I)


# 算账也要有文字证据：模型会看着营业厅背景脑补「屏幕显示电信业务界面」、把座机当「计算器」（实测两张都是）。
# 本机 OCR 在【水印以外】读到办理类的字（手机界面）或数字 / 算式（计算器、手写）才算，否则退回拦截。
# 拦截和算账都是有效打卡、都不填色，退错了只影响格子里的字
_WATERMARK_LINE = re.compile(r"时\s*间|地\s*点|支\s*局|小组成员|拍摄人|工作内容|工作单位|备\s*注|防伪|验真|今日水印|"
                             r"真实可验|中国电信|CHINA\s*TELECOM|活动|经纬度|海拔|"
                             # OCR 常把「时 间」拆成两行，剩下「间:2026.09.2115:38」—— 带日期的行都当水印
                             r"^\s*间|\d{4}\s*[.\-/年]\s*\d{1,2}\s*[.\-/月]\s*\d{1,2}")
_SCREEN_WORDS = re.compile(r"受理|办理|订单|下单|提交|确认|套餐|资费|业务|元/月|月租")
_CALC = re.compile(r"\d{2,}|[+×÷=＝]")
# 手写字 OCR 常读不出来：模型说是拿笔在纸上写画的也认（09-28 #93 拿笔在宣传单上圈画，同事标算账）。
# 不认「算」「计算器」这类字：座机被脑补成计算器的就是这么说的
_HANDWRITING = re.compile(r"笔|手写|纸上|写|圈|画")
# 09-29 放宽：模型证据里写了摊位类的东西，地点没报「摊位」也当摊位（#19「站在摊位前…展板」、#61「红桌摊位」被退成无效）
_STALL_WORDS = re.compile(r"摊位|展台|展板|帐篷|易拉宝|折叠桌")
# 09-29 放宽：拿笔在纸上写 / 圈 / 画 / 算，模型判成板夹或拦截的改判算账（#93「手持笔在翻开的资费单上书写」）
_PEN_WRITING = re.compile(r"笔.{0,15}(?:写|圈|画|算)|手写")


def _scene_text(items_text: list[str]) -> str:
    """去掉水印那几行，剩下的是画面里拍到的字。"""
    return " ".join(t for t in items_text if not _WATERMARK_LINE.search(t))


_INDOOR = ("入户", "办公室")


def refine(cls: str, evidence: str, lines: list[str], place: str = "", folder: bool = False) -> tuple[str, str]:
    """视觉模型的结果再用本机 OCR 核一遍。lines 是 OCR 读到的每一行，place 是模型说的拍摄地点，
    folder 是模型说画面里有板夹 / 资费单摆给客户看。"""
    scene = _scene_text(lines)
    if cls == "异网" and not _OTHER_CARRIER.search(scene):
        cls, evidence = "算账", (evidence + "；屏幕上没认出其他运营商")[:40]
    if cls == "算账" and not (_SCREEN_WORDS.search(scene) or _CALC.search(scene) or _HANDWRITING.search(evidence)):
        # 没有算账的证据：有板夹摆给客户看按板夹（09-28 #56 板夹 + 手机一起讲）；
        # 户外 / 摊位 / 营业厅按拦截；入户 / 办公室只是有人在场 = 无效
        if folder:
            return "板夹", (evidence + "；没认出办理界面，有板夹按板夹")[:40]
        if place in _INDOOR:
            return "无效", (evidence + "；没认出办理界面或算式，入户或办公室无展示按无效")[:40]
        return "拦截", (evidence + "；画面里没认出办理界面或算式，按拦截")[:40]
    return cls, evidence


# 这些行在送模型前涂黑：证件、验证码、病历、长串数字
_SENSITIVE = re.compile(r"身份|公民|居民|住址|出生|民族|签发|有效期|验证码|住院|诊断|出院|病|姓名|\d{5,}")


def mask_and_encode(shrunk: bytes, items: list, max_side: int) -> bytes:
    """把敏感行涂黑，再缩到 max_side 以内，返回 JPEG。"""
    from PIL import Image, ImageDraw

    img = Image.open(io.BytesIO(shrunk)).convert("RGB")
    draw = ImageDraw.Draw(img)
    for box, text in items:
        if not _SENSITIVE.search(text):
            continue
        try:
            xs = [float(p[0]) for p in box]
            ys = [float(p[1]) for p in box]
        except (TypeError, ValueError, IndexError):
            continue
        draw.rectangle([min(xs) - 4, min(ys) - 4, max(xs) + 4, max(ys) + 4], fill=(0, 0, 0))
    img.thumbnail((max_side, max_side))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    return buf.getvalue()


def masked_jpeg(blob: bytes, ocr_settings, max_side: int = 1024) -> bytes:
    """原图 -> 送模型的那个版本（证件 / 验证码 / 病历 / 长数字行涂黑，缩到 max_side）。样本也只存这个。"""
    shrunk, items = _ENGINE.read_items(blob, ocr_settings)
    return mask_and_encode(shrunk, items, max_side)


class VisionJudge:
    DOWN_HOURS = 6       # 额度用完的模型歇多久再试（充值 / 换 key 后能自己回到主模型）

    def __init__(self, s: CheckinSettings) -> None:
        self.s = s
        self._session = requests.Session()
        self._lock = threading.Lock()
        self._window, self._calls = 0.0, 0
        self._down: dict[str, float] = {}      # 模型 -> 到什么时候之前别再试（额度用完 / 没开通）
        self._alerted: dict[str, float] = {}   # 提醒种类 -> 上次提醒时刻，同一种 6 小时最多一次
        self.notify = None                     # 给管理员发单聊的函数，pipeline 装上
        self.last_model = ""
        self.last_error = ""                   # ok / nokey / budget / auth / quota_all / fail

    def _alert(self, kind: str, text: str) -> None:
        now = time.monotonic()
        with self._lock:
            if now - self._alerted.get(kind, -1e9) < self.DOWN_HOURS * 3600:
                return
            self._alerted[kind] = now
        log.warning("视觉模型提醒：%s", text)
        if self.notify is not None:
            try:
                self.notify(text)
            except Exception:
                log.exception("给管理员发视觉模型提醒失败")

    def models(self) -> list[str]:
        """主模型 + 备用模型，去掉正在歇着的。"""
        now = time.monotonic()
        out = []
        for m in [self.s.vision_model, *(self.s.vision_fallback_models or ())]:
            m = str(m).strip()
            if m and m not in out and self._down.get(m, 0) <= now:
                out.append(m)
        return out

    def _budget(self) -> bool:
        limit = self.s.vision_max_calls_per_minute
        if limit <= 0:
            return True
        now = time.monotonic()
        with self._lock:
            if now - self._window >= 60:
                self._window, self._calls = now, 0
            if self._calls >= limit:
                return False
            self._calls += 1
            return True

    def judge(self, jpeg: bytes) -> tuple[str, str, str, bool] | None:
        """
        返回 (类别, 证据, 地点, 有没有板夹摆给客户看)；判不了返回 None（原因在 last_error）。
        主模型额度用完 / 没开通就按顺序换备用模型，并单聊提醒管理员；全用完了也提醒。
        """
        key = os.getenv(self.s.vision_api_key_env, "").strip()
        if not key:
            self.last_error = "nokey"
            self._alert("nokey", f"没配视觉模型的 key（{self.s.vision_api_key_env}），打卡先记「待判定」，配好后会自动补判")
            return None
        if not self._budget():
            self.last_error = "budget"
            log.warning("视觉模型每分钟 %d 次已满，这张记「待判定」", self.s.vision_max_calls_per_minute)
            return None
        primary = str(self.s.vision_model).strip()
        tried = []
        for model in self.models():
            tried.append(model)
            kind, got = self._call(model, key, jpeg)
            if kind == "ok":
                self.last_model, self.last_error = model, "ok"
                if model != primary:
                    self._alert(f"switch:{model}", f"视觉模型 {primary} 额度用完或没开通，已自动换成 {model} 继续判打卡")
                return got
            if kind == "quota":
                self._down[model] = time.monotonic() + self.DOWN_HOURS * 3600
                log.warning("视觉模型 %s 额度用完 / 没开通，%d 小时内换别的模型：%s", model, self.DOWN_HOURS, got)
                continue
            if kind == "auth":
                self.last_error = "auth"
                self._alert("auth", f"视觉模型的 key 不能用了（{got}），打卡先记「待判定」，换好 key 重启后会自动补判")
                return None
            self.last_error = "fail"            # 网络抖动、返回格式不对：这张先待判定，过会儿补判
            return None
        self.last_error = "quota_all"
        self._alert("quota_all", "视觉模型额度全用完了（" + "、".join(
            [primary, *[str(m) for m in (self.s.vision_fallback_models or ())]]) +
            "），打卡先记「待判定」。换 key 或者在百炼关掉「只用免费额度」后会自动补判")
        return None

    _RATE_LIMIT = re.compile(r"Throttling|RateQuota|rate limit", re.I)
    _NO_QUOTA = re.compile(r"quota|FreeTier|Arrearage|access_denied|AccessDenied|not.?exist|Unpurchased", re.I)

    def _call(self, model: str, key: str, jpeg: bytes):
        """调一个模型。返回 ("ok", 结果) / ("quota", 原因) / ("auth", 原因) / ("fail", 原因)。"""
        body = {"model": model, "temperature": 0, "max_tokens": 120,
                "messages": [{"role": "user", "content": [
                    {"type": "image_url",
                     "image_url": {"url": "data:image/jpeg;base64," + base64.b64encode(jpeg).decode()}},
                    {"type": "text", "text": PROMPT}]}]}
        why = ""
        for attempt in (1, 2):
            try:
                r = self._session.post(f"{self.s.vision_base_url.rstrip('/')}/chat/completions",
                                       headers={"Authorization": f"Bearer {key}"},
                                       json=body, timeout=self.s.vision_timeout)
                if r.status_code == 401:
                    return "auth", f"HTTP 401 {r.text[:80]}"
                if r.status_code in (400, 403, 404) and not self._RATE_LIMIT.search(r.text) \
                        and self._NO_QUOTA.search(r.text):
                    return "quota", f"HTTP {r.status_code} {r.text[:120]}"
                if r.status_code != 200:
                    raise RuntimeError(f"HTTP {r.status_code}: {r.text[:160]}")
                txt = r.json()["choices"][0]["message"]["content"]
                return "ok", self._interpret(json.loads(txt[txt.find("{"): txt.rfind("}") + 1]))
            except Exception as exc:
                why = str(exc)[:160]
                log.warning("视觉模型 %s 第 %d 次调用失败：%s", model, attempt, exc)
        return "fail", why

    @staticmethod
    def _interpret(data: dict) -> tuple[str, str, str, bool]:
        """模型的 JSON -> (类别, 证据, 地点, 有没有板夹)，带上本机的一致性兜底。"""
        label = str(data.get("label") or "").strip()
        if label not in DISPLAY:
            raise RuntimeError(f"类别不认识：{label!r}")
        # 自相矛盾时按「有没有人」这一步为准：说有人却判无效 -> 拦截；说没人却判了别的 -> 无效
        people = data.get("people")
        if people is True and label == "无效":
            label = "拦截"
        place = str(data.get("place") or "").strip()
        tv = str(data.get("tv") or "").strip()
        if people is False and label in ("拦截", "板夹", "设备演示") and not (
                label == "设备演示" and tv == "云盒菜单"):
            # 这三类得有客户在；异网 / 算账可以是手机屏幕特写，没人也算；
            # 云盒菜单特写没人也算演示（09-28 #99，用户：拍到云盒菜单就算）
            label = "无效"
        installing = data.get("installing") is True or str(data.get("installing")).lower() == "true"
        folder = data.get("folder") is True or str(data.get("folder")).lower() == "true"
        evidence = str(data.get("evidence") or "")[:40]
        if "监控" in evidence:
            tv = "普通节目"     # 监控多格画面会被认成云盒菜单（09-28 #75）
        talking = not (data.get("talking") is False or str(data.get("talking")).lower() == "false")
        if place != "摊位" and _STALL_WORDS.search(evidence):
            place = "摊位"
        if label in ("板夹", "拦截") and people is not False and _PEN_WRITING.search(evidence):
            label = "算账"
        if place == "摊位":
            # 摊位上一律算拦截：桌上摆着资料设备也是、没拍到正在说话也是（09-24 同事的标法）
            if label in ("板夹", "设备演示", "无效") and people is not False:
                label = "拦截"
            return label, evidence, place, folder
        # 摊位以外，有人但没有有效展示时：入户 / 办公室 = 无效；别处有人交谈 = 拦截，没人交谈 = 无效
        # （09-24 用户：有人说话就算拦截；入户客户在场没展示算无效）
        fallback = "拦截" if (talking and place not in _INDOOR) else "无效"
        if label == "设备演示" and (installing or tv == "普通节目"):
            label = fallback        # 装机接线、电视播普通节目都不是设备演示
        elif label == "拦截" and (place in _INDOOR or not talking):
            label = "无效"
        # refine 里「算账没证据」按同一口径退：没人交谈就当入户处理（退无效）
        return label, evidence, (place if talking else "入户"), folder


# ---------------- 存储 ----------------

class CheckinStore:
    """打卡记录，整份存 JSON。只有机器人进程写；出表 / MCP 每次重新读。"""

    def __init__(self, path: Path) -> None:
        self.path = path
        self._lock = threading.Lock()
        self.records: list[Checkin] = load_records(path)
        self._photos = {r.photo for r in self.records if r.photo}
        self._keys = {dedupe_key(r) for r in self.records}

    def has_photo(self, photo: str) -> bool:
        with self._lock:
            return photo in self._photos

    def has_key(self, key: str) -> bool:
        with self._lock:
            return key in self._keys

    def add(self, rec: Checkin) -> bool:
        key = dedupe_key(rec)
        with self._lock:
            if key in self._keys or (rec.photo and rec.photo in self._photos):
                return False
            self.records.append(rec)
            self._photos.add(rec.photo)
            self._keys.add(key)
            payload = {"records": [asdict(r) for r in self.records]}
        self._write(payload)
        return True

    def find(self, photo: str) -> Checkin | None:
        with self._lock:
            return next((r for r in self.records if r.photo == photo), None)

    def update(self, photo: str, **fields) -> bool:
        """补判用：按照片指纹改一条记录。"""
        with self._lock:
            rec = next((r for r in self.records if r.photo == photo), None)
            if rec is None:
                return False
            for k, v in fields.items():
                setattr(rec, k, v)
            payload = {"records": [asdict(r) for r in self.records]}
        self._write(payload)
        return True

    def replace(self, old_photo: str, new: "Checkin") -> "Checkin | None":
        """
        无效打卡兜底：新的有效照片替换旧的那条无效。留旧的 时间 / 时段（09-29 用户定：用旧的时间），
        其余（类型、颜色、来源、证据、照片指纹、防伪码……）换成新的。返回替换后的记录；旧的没了返回 None。
        """
        with self._lock:
            rec = next((r for r in self.records if r.photo == old_photo), None)
            if rec is None:
                return None
            self._photos.discard(rec.photo)
            self._keys.discard(dedupe_key(rec))
            keep_at, keep_slot = rec.at, rec.slot
            for f in Checkin.__dataclass_fields__:
                setattr(rec, f, getattr(new, f))
            rec.at, rec.slot = keep_at, keep_slot
            self._photos.add(rec.photo)
            self._keys.add(dedupe_key(rec))
            payload = {"records": [asdict(r) for r in self.records]}
        self._write(payload)
        return rec

    def _write(self, payload: dict) -> None:
        tmp = self.path.with_suffix(".tmp")
        for attempt in (1, 2):
            try:
                tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
                keep_version(self.path)          # 写之前留底
                tmp.replace(self.path)
                break
            except Exception as exc:
                if attempt == 2:
                    log.error("打卡记录写不进去（内存里还在）：%s", exc)
                time.sleep(1)


def dedupe_key(rec) -> str:
    """同一张照片的判定键：防伪码优先；认不出就用 水印时间 + 成员 + 工作内容。rec 可以是 Checkin 或水印 dict。"""
    get = rec.get if isinstance(rec, dict) else (lambda k, d="": getattr(rec, k, d))
    code = get("code", "")
    if code:
        return "code:" + code
    at = get("stamp", "") or get("at", "")
    squash = lambda v: re.sub(r"[\s、,，]", "", v or "")
    return f"wm:{at}|{squash(get('members', ''))}|{squash(get('work', ''))}"


def load_records(path: Path) -> list[Checkin]:
    if not path.exists():
        return []
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
        rows = raw.get("records") if isinstance(raw, dict) else None
        keys = Checkin.__dataclass_fields__
        return [Checkin(**{k: v for k, v in r.items() if k in keys}) for r in rows or () if isinstance(r, dict)]
    except Exception as exc:
        log.warning("打卡记录读不了，当作空的：%s", exc)
        return []


# ---------------- 记录一张照片 ----------------

class CheckinRecorder:
    def __init__(self, settings_getter, ocr_settings_getter) -> None:
        self._settings = settings_getter
        self._ocr = ocr_settings_getter
        s = settings_getter()
        self.store = CheckinStore(DATA / s.store_file)
        self._judge = VisionJudge(s)
        self.pending_dir = DATA / "checkin_pending"
        self._retry_lock = threading.Lock()
        self._pool = None
        # 多群（09-29）：identify(发送人账号, 水印小组成员) -> ("ok", 区域) / ("non_activity", None) / ("unknown", None)；
        # 没装就照旧谁发的都记。on_photo(照片, 水印) 每张有水印时间的照片都调一次（支局长人脸点检用）
        self.identify = None
        self.on_photo = None
        # 无效打卡兜底：发送人 -> [(截止时刻, 那条无效的照片指纹)]；有效的来了替换最近那条（只在内存，重启就没了）
        self.retry_minutes_invalid = 10
        self._invalid_windows: dict[str, list[tuple[float, str]]] = {}
        self._win_lock = threading.Lock()

    def open_window(self, sender: str, photo: str) -> float:
        """判成无效 / 待判定：给这个人开 10 分钟重拍窗口，返回截止时刻（给回执写「14:50 前」）。"""
        deadline = time.time() + self.retry_minutes_invalid * 60
        if sender:
            with self._win_lock:
                self._invalid_windows.setdefault(sender, []).append((deadline, photo))
        return deadline

    def _take_window(self, sender: str) -> str | None:
        """有效的来了：取这个人最近一条还在窗口里、而且还是无效 / 待判定的那条（补判改成有效了的跳过）。"""
        if not sender:
            return None
        now = time.time()
        with self._win_lock:
            wins = [w for w in self._invalid_windows.get(sender, []) if w[0] >= now]
            self._invalid_windows[sender] = wins
            while wins:
                _deadline, photo = wins.pop()
                old = self.store.find(photo)
                if old is not None and (old.source == "待判定" or old.label == DISPLAY["无效"]):
                    return photo
        return None

    # ---------- 「待判定」补判（09-28） ----------
    # 视觉模型判不了（额度用完、key 失效、网络）的照片，把【涂黑缩小后】送模型的那张图和
    # 本机 OCR 读到的画面文字（敏感行去掉）存到 data/checkin_pending/，定时再判；
    # 判出来就改记录、删图；留满 pending_keep_days 天还判不了就删图放弃（记录保持待判定）

    @property
    def notify(self):
        return self._judge.notify

    @notify.setter
    def notify(self, fn) -> None:
        self._judge.notify = fn

    def _save_pending(self, photo: str, jpeg: bytes, lines: list[str], work: str) -> None:
        try:
            self.pending_dir.mkdir(parents=True, exist_ok=True)
            (self.pending_dir / f"{photo}.jpg").write_bytes(jpeg)
            meta = {"work": work, "lines": [t for t in lines if not _SENSITIVE.search(t)],
                    "created": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
            (self.pending_dir / f"{photo}.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
        except Exception:
            log.exception("待判定照片没存下来（这张以后补判不了）")

    def _drop_pending(self, photo: str) -> None:
        for ext in (".jpg", ".json"):
            try:
                (self.pending_dir / f"{photo}{ext}").unlink()
            except FileNotFoundError:
                pass

    def pending_count(self) -> int:
        return len(list(self.pending_dir.glob("*.json"))) if self.pending_dir.is_dir() else 0

    def retry_pending(self) -> dict:
        """把存着的「待判定」再判一遍。返回 {done, left, dropped, reason}。同一时间只跑一个。"""
        out = {"done": 0, "left": 0, "dropped": 0, "reason": ""}
        if not self._retry_lock.acquire(blocking=False):
            out["left"] = self.pending_count()
            out["reason"] = "正在补判"
            return out
        try:
            s = self._settings()
            self._judge.s = s
            cutoff = datetime.now() - timedelta(days=s.pending_keep_days)
            for meta_path in sorted(self.pending_dir.glob("*.json")) if self.pending_dir.is_dir() else ():
                photo = meta_path.stem
                rec = self.store.find(photo)
                jpg = meta_path.with_suffix(".jpg")
                if rec is None or rec.source != "待判定" or not jpg.is_file():
                    self._drop_pending(photo)
                    continue
                try:
                    meta = json.loads(meta_path.read_text(encoding="utf-8"))
                    created = datetime.strptime(meta.get("created", ""), "%Y-%m-%d %H:%M:%S")
                except Exception:
                    meta, created = {}, datetime.now()
                if created < cutoff:
                    log.warning("打卡 %s（%s）%d 天都没判出来，放弃补判，照片已删", rec.at, rec.branch, s.pending_keep_days)
                    self._drop_pending(photo)
                    out["dropped"] += 1
                    continue
                label, color, source, evidence = self._vision(meta.get("work", ""), jpg.read_bytes(),
                                                              meta.get("lines") or [])
                if source == "待判定":
                    if self._judge.last_error in ("nokey", "auth", "quota_all", "budget"):
                        out["reason"] = self._judge.last_error
                        break           # 整体判不了（没额度 / key 坏 / 限流），别一张张白试
                    continue
                self.store.update(photo, label=label, color=color, source=source, evidence=evidence,
                                  model=self._judge.last_model)
                self._drop_pending(photo)
                out["done"] += 1
                log.info("补判打卡 | %s %s %s | %s（%s）", rec.at, rec.branch, rec.members, label, evidence)
        finally:
            self._retry_lock.release()
        out["left"] = self.pending_count()
        if out["done"] or out["dropped"]:
            log.info("补判打卡：判出 %d 张，放弃 %d 张，还剩 %d 张", out["done"], out["dropped"], out["left"])
        return out

    def start_retry(self, pool) -> None:
        """后台每 retry_minutes 分钟补判一次（交给打卡的单线程池，和新照片排队，不会同时写记录）。"""
        self._pool = pool

        def loop() -> None:
            while True:
                time.sleep(max(1, int(self._settings().retry_minutes)) * 60)
                try:
                    if self.pending_count():
                        pool.submit(self.retry_pending)
                except Exception:
                    log.exception("提交补判出错")

        threading.Thread(target=loop, name="checkin-retry", daemon=True).start()

    def retry_now(self, timeout: float = 300) -> dict:
        """马上补判并等结果（打卡通报前、管理员「补判打卡」）。没有待判定就直接返回。"""
        if not self.pending_count():
            return {"done": 0, "left": 0, "dropped": 0, "reason": ""}
        if self._pool is None:
            return self.retry_pending()
        return self._pool.submit(self.retry_pending).result(timeout=timeout)

    def record(self, msg_id: str, blobs: list[bytes], sender: str = "") -> list[tuple[str, object]]:
        """
        每张照片记一次打卡。返回每张的结果（单聊回执用）：
          ("recorded", Checkin) / ("dup", 水印时间) / ("not_watermark", None) / ("no_time", None) / ("error", None)
        """
        s = self._settings()
        if not s.enabled:
            return []
        self._judge.s = s
        out = []
        for blob in blobs:
            try:
                out.append(self._one(s, msg_id, blob, sender))
            except Exception:
                log.exception("打卡记录出错（%s）", msg_id)
                out.append(("error", None))
        return out

    def read(self, blob: bytes):
        """OCR 读水印。不是今日水印照片返回 None；否则 (水印 dict, 缩过的图, OCR 行)。"""
        shrunk, items = _ENGINE.read_items(blob, self._ocr())
        text = " ".join(t for _b, t in items)
        if not is_watermark(text):
            return None
        return parse_watermark(text), shrunk, items

    def classify(self, wm: dict, shrunk: bytes, items: list) -> tuple[str, str, str, str]:
        """(格子里的字, 颜色, 来源, 证据)。业务 / 原样照写不调模型。"""
        s = self._settings()
        self._judge.s = s
        ruled = rule_label(wm["work"], s.keep_as_is)
        if ruled:
            return (*ruled, "")
        return self._vision(wm["work"], mask_and_encode(shrunk, items, s.image_max_side), [t for _b, t in items])

    def _vision(self, work: str, jpeg: bytes, lines: list[str]) -> tuple[str, str, str, str]:
        """视觉模型判 + 本机兜底。jpeg 是已经涂黑缩小的图，lines 是 OCR 读到的每一行。"""
        got = self._judge.judge(jpeg)
        if not got:
            return "待判定", "", "待判定", ""
        cls, evidence, place, folder = got
        if place == "店铺" and "上门" in work and "拦截" not in work:
            # 上门去店里找老板 = 入户：只聊天没展示算无效（09-28 用户定，#77 #79 #102）；
            # 店里拦路人（工作内容写拦截）不受影响
            place = "入户"
            if cls == "拦截":
                cls, evidence = "无效", (evidence + "；上门进店没展示东西按无效")[:40]
        cls, evidence = refine(cls, evidence, lines, place=place, folder=folder)
        return DISPLAY[cls], (RED if cls == "无效" else ""), "视觉", evidence

    def judge_blob(self, blob: bytes) -> tuple[str, str] | None:
        """样本回执用：只判不记。返回 (格子里的字, 证据或来源说明)；不是水印照片返回 None。"""
        got = self.read(blob)
        if got is None:
            return None
        label, _color, source, evidence = self.classify(*got)
        return label, evidence or ("工作内容里有业务，按业务算" if source == "业务" else source)

    def _one(self, s: CheckinSettings, msg_id: str, blob: bytes, sender: str = "") -> tuple[str, object]:
        photo = hashlib.sha1(blob).hexdigest()[:16]
        if self.store.has_photo(photo):
            log.info("打卡照片重复发送，不重复记（%s）", photo)
            old = self.store.find(photo)
            return "dup", (old.at if old else "")
        got = self.read(blob)
        if got is None:
            return "not_watermark", None
        wm, shrunk, items = got
        if not wm["stamp"]:
            log.info("水印照片没认出时间，不记打卡")
            return "no_time", None
        if self.on_photo is not None:
            try:                        # 支局长人脸点检：谁发的、记不记打卡都要看（09-29 用户定）
                self.on_photo(blob, wm, photo)
            except Exception:
                log.exception("支局长人脸点检出错（不影响打卡）")
        # 企微转发会重压缩，文件指纹对不上 —— 按防伪码再判一次，省得重复调模型
        if self.store.has_key(dedupe_key(wm)):
            log.info("打卡照片重复发送，不重复记（%s）", wm["code"] or wm["stamp"])
            return "dup", wm["stamp"]

        region = ""
        if self.identify is not None:
            # 只有活动人员打卡（09-29）：① 发送人账号 ② 认不出再看水印小组成员 ③ 还不行就不记，回一句让他改信息重发
            status, region = self.identify(sender, wm.get("members", ""), wm.get("stamp", ""), wm.get("branch", ""))
            if status != "ok":
                log.info("打卡没记（%s）| 发送人=%s | 小组成员=%s", status, sender or "?", wm.get("members") or "空")
                return status, None
            region = region or ""
        label, color, source, evidence = self.classify(wm, shrunk, items)
        rec = Checkin(at=wm["stamp"], slot=slot_of(wm["stamp"][11:16], s.slots),
                      branch=wm["branch"], members=wm["members"], work=wm["work"],
                      label=label, color=color, source=source, evidence=evidence,
                      photo=photo, msg_id=msg_id,
                      received=datetime.now().strftime("%Y-%m-%d %H:%M:%S"), sender=sender,
                      code=wm["code"], model=self._judge.last_model if source == "视觉" else "",
                      region=region)
        if source == "待判定":
            # 存涂黑缩小后的图等补判（09-28 用户同意：最多留 pending_keep_days 天，判出来就删）
            self._save_pending(photo, mask_and_encode(shrunk, items, s.image_max_side),
                               [t for _b, t in items], wm["work"])
        invalid = source == "待判定" or label == DISPLAY["无效"]
        if not invalid:
            old_photo = self._take_window(sender)
            if old_photo:
                old = self.store.find(old_photo)
                old_at = old.at if old else ""
                replaced = self.store.replace(old_photo, rec)
                if replaced is not None:
                    self._drop_pending(old_photo)          # 旧的要是待判定，存的图不用再补判了
                    log.info("打卡替换无效 | %s | 发送人=%s | %s -> %s（%s）", old_at, sender or "?",
                             DISPLAY["无效"], rec.label, rec.source)
                    return "replaced", (replaced, old_at)
        if not self.store.add(rec):
            return "dup", rec.at
        log.info("打卡 | %s %s | 发送人=%s | %s %s | 工作内容=%s | %s（%s）%s", rec.at, rec.slot or "时段外",
                 sender or "?", rec.branch or "?", rec.members or "?", rec.work or "空", rec.label,
                 rec.source, f" 证据：{evidence}" if evidence else "")
        if invalid:
            return "invalid", (rec, self.open_window(sender, photo))
        return "recorded", rec


# ---------------- 打卡人 ----------------
# 业务定的（09-24）：谁发的打卡照片，就算到谁所在的那一组 —— 杨桐发的算杨桐那组。
# 长连接只给发送人的账号 ID（如 HuangTao），不给姓名，所以要把 ID 对到名册里的人：
#   ① data/userid_map.json 手填的对照 {"账号ID": "姓名"} 优先
#   ② 账号 ID 是姓名拼音的（企业默认就是），按拼音自动对；对上唯一一个人才算
#   多音字姓氏（曾 / 单 …）几种读音都试；同拼音两个人就不猜，要手填

USERID_MAP_FILE = "userid_map.json"
_map_cache: dict = {"mtime": None, "map": {}}
_py_cache: dict = {"names": None, "index": {}}


def _userid_map() -> dict:
    p = DATA / USERID_MAP_FILE
    try:
        mtime = p.stat().st_mtime
    except OSError:
        return {}
    if _map_cache["mtime"] != mtime:
        try:
            raw = json.loads(p.read_text(encoding="utf-8"))
            _map_cache["map"] = {str(k): str(v).strip() for k, v in (raw or {}).items()
                                 if not str(k).startswith("_") and str(v).strip()}
        except Exception as exc:
            log.warning("%s 读不了：%s", USERID_MAP_FILE, exc)
            _map_cache["map"] = {}
        _map_cache["mtime"] = mtime
    return _map_cache["map"]


def _pinyin_keys(name: str) -> set[str]:
    from pypinyin import Style, pinyin
    combos = [""]
    for opts in pinyin(name, style=Style.NORMAL, heteronym=True):
        combos = [c + o for c in combos for o in opts[:3]][:32]
    return {c.lower() for c in combos}


def sender_name(userid: str, roster) -> str:
    """发送人账号 ID -> 名册里的姓名；认不出返回空串。"""
    if not userid:
        return ""
    manual = _userid_map().get(userid)
    if manual:
        return manual
    key = re.sub(r"[^a-z]", "", userid.lower())
    if not key:
        return ""
    names = tuple(sorted(roster.people))
    if _py_cache["names"] != names:
        index: dict[str, set] = {}
        for n in names:
            try:
                for k in _pinyin_keys(n):
                    index.setdefault(k, set()).add(n)
            except Exception:
                continue
        _py_cache.update(names=names, index=index)
    hits = _py_cache["index"].get(key, set())
    return next(iter(hits)) if len(hits) == 1 else ""


# ---------------- 归组 ----------------

def resolve(rec: Checkin, roster) -> tuple[str, str]:
    """
    按当前名册把一条打卡归到 (战队, 小组成员串)。小组认不出返回 (战队, "")。
    顺序：发送人所在的组 > 水印「小组成员」里出现的人所在的组 > 只按支局归战队。
    """
    who = sender_name(rec.sender, roster)
    info = roster.people.get(who) if who else None
    if info:
        return roster.normalize_team(info.get("team", "")), info.get("group", "")
    compact = re.sub(r"[\s、,，/]", "", rec.members or "")
    hits: dict[str, int] = {}
    team_of: dict[str, str] = {}
    for name, info in roster.people.items():
        if name and name in compact:
            g = info.get("group", "")
            hits[g] = hits.get(g, 0) + 1
            team_of[g] = roster.normalize_team(info.get("team", ""))
    if hits:
        g = max(hits, key=hits.get)
        return team_of[g], g
    return roster.normalize_team(rec.branch), ""


def select(records: list[Checkin], day: str = "", slot: str = "",
           since: str = "", until: str = "", slots: list | None = None) -> list[Checkin]:
    """
    按天 / 时段 / 时间范围筛。传了 slots 就按【当前】时段规则重算每条的时段，
    并去掉不统计的（最后一个时段之后）—— 规则改了，旧记录也跟着新规则算。
    """
    out = []
    for r in records:
        if slots is not None:
            r = Checkin(**{**asdict(r), "slot": slot_of(r.at[11:16], slots)})
            if not r.slot:
                continue
        if day and r.at[:10] != day:
            continue
        if slot and r.slot != slot:
            continue
        if since and r.at < since[:16]:
            continue
        if until and r.at > until[:16]:
            continue
        out.append(r)
    return sorted(out, key=lambda r: r.at)


# ---------------- 出表 ----------------

def team_display(team: str) -> str:
    """表里写「通惠战区」，文字通报里写「通惠战队」（通报样表就是这样）。"""
    return team.replace("战队", "战区") if team else "认不出战区"


def table_blocks(records: list[Checkin], roster, columns: int = 13):
    """
    打卡表的数据：(列数, [(战队, 户数, [(小组号, 成员, [打卡…]), …]), …])。
    战队、小组按名册顺序；水印认不出小组的放在该战队最后一行「未识别」。Excel 和通报图片都用它。
    """
    by_group: dict[tuple[str, str], list[Checkin]] = {}
    for r in records:
        by_group.setdefault(resolve(r, roster), []).append(r)
    width = max([columns] + [len(v) for v in by_group.values()])
    teams = list(roster.teams)
    extra = sorted({t for t, _g in by_group} - set(teams) - {""})
    blocks = []
    for team in teams + extra + ([""] if any(t == "" for t, _g in by_group) else []):
        groups = [(g.label, g.members) for g in roster.groups() if g.team == team]
        if any(t == team and g == "" for t, g in by_group):
            groups.append(("未识别", ""))
        if not groups:
            continue
        rows = [(label, members, sorted(by_group.get((team, members), []), key=lambda r: r.at))
                for label, members in groups]
        blocks.append((team, sum(len(r[2]) for r in rows), rows))
    return width, blocks


def write_checkin_sheet(ws, records: list[Checkin], roster, title: str, columns: int = 13) -> None:
    """照原表「打卡」sheet 的样子画。records 已经按天 + 时段筛好。"""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    fills = {GREEN: PatternFill("solid", fgColor="92D050"), RED: PatternFill("solid", fgColor="FF0000")}
    bold = Font(bold=True)

    width, blocks = table_blocks(records, roster, columns)
    last_col = 5 + width

    # 表头三行
    ws.merge_cells(start_row=1, start_column=1, end_row=3, end_column=1)
    ws.cell(1, 1, "触点总量")
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=last_col)
    ws.cell(1, 2, title).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=last_col)
    ws.cell(2, 2, "上门时间")
    for c, h in enumerate(["战区", "小组", "成员", "合计"], start=2):
        ws.cell(3, c, h)
    for i in range(width):
        ws.cell(3, 6 + i, f"第{i + 1}户")
    for r in range(1, 4):
        for c in range(1, last_col + 1):
            ws.cell(r, c).border, ws.cell(r, c).alignment = border, center
            if r == 3 or c == 1:
                ws.cell(r, c).font = bold

    row = 4
    for team, total, rows in blocks:
        top = row
        for label, members, recs in rows:
            ws.cell(row, 3, label)
            mc = ws.cell(row, 4, members or "（水印成员认不出）")
            if not recs and members:
                mc.fill = fills[RED]          # 一条打卡都没有
            ws.cell(row, 5, len(recs))
            for i, rec in enumerate(recs):
                cell = ws.cell(row, 6 + i, f"{rec.at[11:16]}（{rec.label}）")
                if rec.color in fills:
                    cell.fill = fills[rec.color]
            for c in range(1, last_col + 1):
                ws.cell(row, c).border, ws.cell(row, c).alignment = border, center
            row += 1
        ws.merge_cells(start_row=top, start_column=1, end_row=row - 1, end_column=1)
        ws.merge_cells(start_row=top, start_column=2, end_row=row - 1, end_column=2)
        ws.cell(top, 1, total).font = bold
        ws.cell(top, 2, team_display(team)).font = bold

    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 11
    ws.column_dimensions["C"].width = 7
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 6
    for i in range(width):
        ws.column_dimensions[get_column_letter(6 + i)].width = 17
    ws.freeze_panes = "F4"


def span_label(s: CheckinSettings, day: str) -> str:
    """标题里的日期：配了活动日期（9.24-9.27）就用它，否则写当天（9.24）。"""
    if s.span:
        return str(s.span)
    d = datetime.strptime(day, "%Y-%m-%d")
    return f"{d.month}.{d.day}"


def checkin_title(s: CheckinSettings, day: str, campaign_title: str, start: str, end: str) -> str:
    """「9.24-9.27日城区营销活动过程管控情况（13:00-15:30时段）」。"""
    return f"{span_label(s, day)}日{campaign_title or '营销活动'}过程管控情况（{start}-{end}时段）"


def of_region(records: list[Checkin], region: str | None) -> list[Checkin]:
    """多群：只要某个区的打卡。region 为 None 不筛；老记录没有 region 算城区。"""
    if region is None:
        return records
    return [r for r in records if (r.region or "城区") == region]


def day_sheets(s: CheckinSettings, day: str, campaign_title: str,
               region: str | None = None) -> list[tuple[str, str, list[Checkin]]]:
    """某天按时段的几张表：[(sheet 名, 标题, 记录)]。day = YYYY-MM-DD。region = 城区 / 农村（多群）。"""
    records = of_region(load_records(DATA / s.store_file), region)
    out = []
    for slot in s.slots:
        name = str(slot.get("name", ""))
        title = checkin_title(s, day, campaign_title, str(slot.get("start")), str(slot.get("end")))
        out.append((f"打卡-{name}", title, select(records, day=day, slot=name, slots=s.slots)))
    return out


def summarize(records: list[Checkin], roster) -> dict:
    """给 MCP / 助手用：各小组业务、有效、无效、合计，外加一条都没有的小组。"""
    rows: dict[tuple[str, str], dict] = {}
    for g in roster.groups():
        rows[(g.team, g.members)] = {"战队": g.team, "小组": g.label, "成员": g.members,
                                     "业务": 0, "有效": 0, "无效": 0, "待判定": 0, "合计": 0, "明细": []}
    for r in records:
        key = resolve(r, roster)
        row = rows.setdefault(key, {"战队": key[0] or "认不出", "小组": "未识别", "成员": r.members,
                                    "业务": 0, "有效": 0, "无效": 0, "待判定": 0, "合计": 0, "明细": []})
        row["合计"] += 1
        if r.color == GREEN:
            row["业务"] += 1
        elif r.color == RED:
            row["无效"] += 1
        elif r.source == "待判定":
            row["待判定"] += 1
        else:
            row["有效"] += 1
        row["明细"].append(f"{r.at[5:16]} {r.slot or '时段外'}（{r.label}）")
    groups = list(rows.values())
    unknown = sorted({r.sender for r in records if r.sender and not sender_name(r.sender, roster)})
    return {
        "认不出的打卡人账号": unknown,
        "打卡总数": len(records),
        "业务": sum(g["业务"] for g in groups),
        "有效": sum(g["有效"] for g in groups),
        "无效": sum(g["无效"] for g in groups),
        "待判定": sum(g["待判定"] for g in groups),
        "没有打卡的小组": [f"{g['战队']} {g['小组']}（{g['成员']}）" for g in groups if g["合计"] == 0],
        "各小组": [g for g in groups if g["合计"]],
    }
