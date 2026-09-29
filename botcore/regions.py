"""
多群（09-29 国庆版）：城区 / 农村两个区域，每区一张【支局名册】（全体人员）+ 一张【活动名册】（参加活动的人），
外加一份【支局长名单】（支局长打卡点检表用）。规则见 docs/国庆多群版功能说明.md 第二节。

    文件                        内容
    roster.json                 城区活动名册（沿用原来那张：姓名 → 分组 / 战区）
    roster_rural.json           农村活动名册（同上格式）
    branch_city.json            城区支局名册：姓名 → 支局（可带「小区营销」标记）
    branch_rural.json           农村支局名册：同上
    chiefs.json                 支局长名单：类别 / 支局 / 支局长

一个人归成一类（按人，不按群）：城区活动 / 农村活动 / 城区非活动 / 农村非活动 / 未识别。
活动名册里的人算活动人员；只在支局名册里的算非活动；哪张都没有 = 未识别。

多群模式什么时候开：config.yaml 的 groups 里配了 rural（农村活动群）或 hq（攻坚群）。
没开就是原来的单群行为（只有城区活动名册那一张）。
"""

from __future__ import annotations

import copy
import io
import json
import logging
import re
from dataclasses import dataclass
from pathlib import Path

from .campaign import NotRosterSheet, Roster, RosterImportError
from .paths import DATA

log = logging.getLogger("botcore.regions")

CITY, RURAL = "城区", "农村"
REGIONS = (CITY, RURAL)
ACTIVITY, OTHER, UNKNOWN = "活动", "非活动", "未识别"

ACTIVITY_FILES = {CITY: "roster.json", RURAL: "roster_rural.json"}
BRANCH_FILES = {CITY: "branch_city.json", RURAL: "branch_rural.json"}
CHIEFS_FILE = "chiefs.json"

# 管理员单聊先发这句，再发 Excel：导哪张
ROSTER_COMMANDS = {
    "城区活动名册": (CITY, ACTIVITY), "农村活动名册": (RURAL, ACTIVITY),
    "城区支局名册": (CITY, "支局"), "农村支局名册": (RURAL, "支局"),
    "支局长名单": ("", "支局长"),
}

_SUFFIX = re.compile(r"(?:支局|战队|战区|分局)$")


def branch_stem(name: str) -> str:
    """「通惠支局」「通惠战队」「通惠」→「通惠」。"""
    return _SUFFIX.sub("", (name or "").strip()).strip()


def multi_enabled(store) -> bool:
    try:
        return any(g.role in ("rural", "hq") for g in store.groups())
    except Exception:
        return False


def group_role(store, chat_id: str) -> str:
    """群号 → 角色 city / rural / hq；不在 groups 里返回空串。"""
    try:
        for g in store.groups():
            if g.chat_id == chat_id:
                return g.role
    except Exception:
        pass
    return ""


def role_chats(store, role: str) -> tuple[str, ...]:
    try:
        return tuple(g.chat_id for g in store.groups() if g.role == role)
    except Exception:
        return ()


REGION_OF_ROLE = {"city": CITY, "rural": RURAL}
DAILY_FILE = "roster_days.json"     # 国庆每天的活动名单（09-29）


# ---------------- 国庆活动计划表 → 每天的名单 ----------------

_DAY_HEAD = re.compile(r"(\d{1,2})\s*月\s*(\d{1,2})\s*日\s*活动人员")
_GROUP_START = re.compile(r"(?=\d+\s*组\s*[：:])")
_GROUP_CELL = re.compile(r"\s*(\d+)\s*组\s*[：:]\s*(.*)", re.S)


def split_groups(cell) -> list[list[str]]:
    """「1组：甲、乙\n2组：丙，丁 3组：戊 己」→ [[甲, 乙], [丙, 丁], [戊, 己]]。分隔符顿号 / 逗号 / 空格都认。"""
    out = []
    for part in _GROUP_START.split(str(cell or "")):
        m = _GROUP_CELL.match(part)
        if m:
            names = [n for n in re.split(r"[、，,\s/]+", m.group(2).strip()) if n]
            if names:
                out.append(names)
    return out


def _day_key(day: str, year: int) -> str:
    """「10.3」「10月3日」「2026-10-03」→「2026-10-03」。"""
    m = re.match(r"\s*(\d{4})-(\d{1,2})-(\d{1,2})", day)
    if m:
        return f"{int(m.group(1))}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    m = re.match(r"\s*(\d{1,2})\s*[.月]\s*(\d{1,2})", day)
    return f"{year}-{int(m.group(1)):02d}-{int(m.group(2)):02d}" if m else day


def parse_daily_plan(data: bytes, region: str, year: int, fixes: dict | None = None) -> tuple[dict, list[str]]:
    """
    国庆活动计划表 → ({日期: {区域: {"teams", "people"}}}, 提醒)。
    表：「支局」一列，每天一列「X月X日活动人员」，格子里「1组：甲、乙\n2组：…」。
    fixes（用户确认过的改正，不进 git，放数据目录）：
        rename    {错字: 正确写法}                                    「胡诚」→「胡城」
        groups    [{day, branch, group, members}]                   某天某支局第几组整组改成谁
        region_of {支局: 区域}                                        新盛在城区表里、算农村
    """
    from openpyxl import load_workbook
    fixes = fixes or {}
    rename = fixes.get("rename") or {}
    region_of = fixes.get("region_of") or {}
    group_fix = {(_day_key(str(g["day"]), year), branch_stem(g["branch"]), int(g["group"])): list(g["members"])
                 for g in fixes.get("groups") or []}
    ws = load_workbook(io.BytesIO(data), data_only=True).worksheets[0]
    head_row, day_cols, branch_col = 0, {}, 0
    for row in ws.iter_rows(min_row=1, max_row=12):
        for c in row:
            v = str(c.value or "").strip()
            m = _DAY_HEAD.search(v)
            if m:
                head_row = c.row
                day_cols[c.column] = f"{year}-{int(m.group(1)):02d}-{int(m.group(2)):02d}"
            elif v == "支局":
                branch_col = c.column
    if not day_cols or not branch_col:
        raise RosterImportError("没认出计划表：要有「支局」一列和「X月X日活动人员」几列")
    out: dict = {}
    notes: list[str] = []
    for r in range(head_row + 1, ws.max_row + 1):
        br = branch_stem(str(ws.cell(r, branch_col).value or ""))
        if not br:
            continue
        reg = region_of.get(br, region)
        team = br + "战队"
        for col, day in sorted(day_cols.items()):
            groups = [[rename.get(n, n) for n in g] for g in split_groups(ws.cell(r, col).value)]
            for i in range(len(groups)):
                if (day, br, i + 1) in group_fix:
                    groups[i] = group_fix[(day, br, i + 1)]
            seen: dict[str, int] = {}
            for i, g in enumerate(groups, 1):
                for n in g:
                    if n in seen:
                        notes.append(f"{day[5:]} {br}：{n} 同时在 {seen[n]} 组和 {i} 组（按 {i} 组记）")
                    seen[n] = i
            if not groups:
                continue
            ros = out.setdefault(day, {}).setdefault(reg, {"teams": [], "people": {}})
            if team not in ros["teams"]:
                ros["teams"].append(team)
            for g in groups:
                for n in g:
                    ros["people"][n] = {"group": "、".join(g), "team": team}
    return out, notes


def merge_daily(parts: list[dict], order: list[str] | None = None) -> dict:
    """几张计划表拼成一份；战队按 order（总表的支局顺序）排，人按战队顺序重排（小组编号跟着对）。"""
    days: dict = {}
    for part in parts:
        for day, regs in part.items():
            for reg, ros in regs.items():
                cur = days.setdefault(day, {}).setdefault(reg, {"teams": [], "people": {}})
                cur["teams"] += [t for t in ros["teams"] if t not in cur["teams"]]
                cur["people"].update(ros["people"])
    rank = {b + "战队": i for i, b in enumerate(order or [])}
    for regs in days.values():
        for ros in regs.values():
            ros["teams"].sort(key=lambda t: rank.get(t, len(rank)))
            ros["people"] = {n: p for t in ros["teams"] for n, p in ros["people"].items() if p["team"] == t}
    return {"days": dict(sorted(days.items()))}
ALL = "全部"
SCOPES = (CITY, RURAL, ALL)


def scoped(camp, store, scope: str | None):
    """
    查数用的台账：城区 / 农村 = 那个区活动人员的单（那个区的名册和目标）；全部 = 所有人的单按支局汇总。
    单群模式（没有四张名册）原样返回。
    """
    if getattr(camp, "rosters", None) is None:
        return camp
    if scope in REGIONS:
        return camp.view(scope, store.campaign_settings().targets_for(scope))
    try:
        branches = [b for bs in store.guoqing_settings().classes.values() for b in bs]
    except Exception:
        branches = []
    return camp.view_all(branches)


def scope_of_chat(store, chat_type: str, chat_id: str) -> str | None:
    """问答：城区群 → 城区，农村群 → 农村，攻坚群 / 单聊 → 全部；单群模式 None。"""
    if not multi_enabled(store):
        return None
    role = group_role(store, chat_id) if chat_type != "p2p" else ""
    return REGION_OF_ROLE.get(role, ALL)


# ---------------- Excel 读表头 ----------------

def _rows(data: bytes):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise RosterImportError(f"文件打不开（要 .xlsx，旧的 .xls 请另存为 xlsx）：{exc}")
    for ws in wb.worksheets:
        rows = [[str(c).strip() if c is not None else "" for c in r] for r in ws.iter_rows(values_only=True)]
        yield ws.title, rows


def _find_header(rows: list[list[str]], want: dict[str, re.Pattern], need: tuple[str, ...]):
    """前 10 行里找表头：want = {字段: 正则}；need 里的字段都找到才算。返回 (表头行号, {字段: 列号})。"""
    for r, cells in enumerate(rows[:10]):
        cols = {}
        for key, pat in want.items():
            for c, v in enumerate(cells):
                if v and pat.fullmatch(v) and key not in cols:
                    cols[key] = c
        if all(k in cols for k in need):
            return r, cols
    return None, None


# ---------------- 支局名册 ----------------

class BranchRoster:
    """姓名 → {branch: 支局, community: 小区 / 场镇营销}。"""

    WANT = {"name": re.compile(r"姓名|人员|员工姓名|营销人员"),
            "branch": re.compile(r"支局|所属支局|单位|部门"),
            "community": re.compile(r".*(?:小区|场镇).*")}

    def __init__(self, path: Path) -> None:
        self.path = Path(path)
        self.people: dict[str, dict] = {}
        self._mtime = None
        self.reload()

    def reload(self) -> None:
        try:
            raw = json.loads(self.path.read_text(encoding="utf-8"))
            self._mtime = self.path.stat().st_mtime
        except FileNotFoundError:
            self.people = {}
            return
        except Exception as exc:
            log.error("支局名册 %s 读不了：%s", self.path.name, exc)
            return
        self.people = dict(raw.get("people") or {})

    def maybe_reload(self) -> bool:
        try:
            m = self.path.stat().st_mtime
        except OSError:
            return False
        if m == self._mtime:
            return False
        self.reload()
        return True

    def import_workbook(self, data: bytes, dry_run: bool = False) -> str:
        people: dict[str, dict] = {}
        found = False
        for _title, rows in _rows(data):
            hr, cols = _find_header(rows, self.WANT, ("name", "branch"))
            if hr is None:
                continue
            found = True
            last_branch = ""
            for cells in rows[hr + 1:]:
                def at(key):
                    c = cols.get(key)
                    return cells[c] if c is not None and c < len(cells) else ""
                name, branch = at("name"), at("branch") or last_branch     # 合并单元格：空的是「同上」
                if branch:
                    last_branch = branch
                if not name or name in ("姓名",):
                    continue
                if not branch:
                    raise RosterImportError(f"「{name}」这行缺支局")
                comm = at("community")
                people[name] = {"branch": branch_stem(branch),
                                "community": bool(comm) and comm not in ("否", "无", "0", "×", "-")}
            break
        if not found:
            raise NotRosterSheet("表里没有「姓名 / 支局」两列的表头")
        if len(people) < 2:
            raise RosterImportError("表头下面没读到人")
        branches = sorted({v["branch"] for v in people.values()})
        if not dry_run:
            self.path.write_text(json.dumps({"people": people}, ensure_ascii=False, indent=1), encoding="utf-8")
            self.reload()
        n_comm = sum(1 for v in people.values() if v["community"])
        return f"{len(people)} 人 / {len(branches)} 个支局" + (f"（小区 / 场镇营销 {n_comm} 人）" if n_comm else "")


# ---------------- 支局长名单 ----------------

class ChiefList:
    """[{class: 一类, branch: 通惠, name: 翁艺丹}]，顺序就是点检表的行顺序。"""

    WANT = {"name": re.compile(r"支局长|姓名"), "branch": re.compile(r"支局"),
            "cls": re.compile(r"类别|分类|类型")}

    def __init__(self, path: Path) -> None:
        self.path = Path(path)
        self.chiefs: list[dict] = []
        self._mtime = None
        self.reload()

    def reload(self) -> None:
        try:
            raw = json.loads(self.path.read_text(encoding="utf-8"))
            self._mtime = self.path.stat().st_mtime
        except FileNotFoundError:
            self.chiefs = []
            return
        except Exception as exc:
            log.error("支局长名单读不了：%s", exc)
            return
        self.chiefs = list(raw.get("chiefs") or [])

    def maybe_reload(self) -> bool:
        try:
            m = self.path.stat().st_mtime
        except OSError:
            return False
        if m == self._mtime:
            return False
        self.reload()
        return True

    def names(self) -> set[str]:
        return {c["name"] for c in self.chiefs}

    def of(self, name: str) -> dict | None:
        return next((c for c in self.chiefs if c["name"] == name), None)

    def import_workbook(self, data: bytes, dry_run: bool = False) -> str:
        chiefs: list[dict] = []
        found = False
        for _title, rows in _rows(data):
            # 「支局长」和「支局」都能匹配「支局」—— 先认「支局长」那列，再认剩下的「支局」
            hr = None
            for r, cells in enumerate(rows[:10]):
                if any(v == "支局长" for v in cells) and any(v == "支局" for v in cells):
                    hr = r
                    break
            if hr is None:
                continue
            found = True
            head = rows[hr]
            c_name = head.index("支局长")
            c_branch = head.index("支局")
            c_cls = next((i for i, v in enumerate(head) if self.WANT["cls"].fullmatch(v)), None)
            last_cls = last_branch = ""
            for cells in rows[hr + 1:]:
                def at(c):
                    return cells[c] if c is not None and c < len(cells) else ""
                name = at(c_name)
                branch = at(c_branch) or last_branch
                cls = at(c_cls) or last_cls
                if branch:
                    last_branch = branch
                if cls:
                    last_cls = cls
                if not name or name == "支局长":
                    continue
                chiefs.append({"class": cls, "branch": branch_stem(branch), "name": name})
            break
        if not found:
            raise NotRosterSheet("表里没有「支局 / 支局长」两列的表头")
        if not chiefs:
            raise RosterImportError("表头下面没读到支局长")
        if not dry_run:
            self.path.write_text(json.dumps({"chiefs": chiefs}, ensure_ascii=False, indent=1), encoding="utf-8")
            self.reload()
        return f"{len(chiefs)} 个支局长 / {len({c['branch'] for c in chiefs})} 个支局"


# ---------------- 认人 ----------------

@dataclass
class Identity:
    name: str = ""
    region: str = ""          # 城区 / 农村 / ""
    role: str = UNKNOWN       # 活动 / 非活动 / 未识别
    team: str = ""            # 活动人员的战队（活动名册里的「战区」）
    group: str = ""           # 活动人员的小组（成员串）
    branch: str = ""          # 支局（简称：通惠）
    community: bool = False   # 小区 / 场镇营销人员（总表里新增宽带、中高套 ×1.1）


class _AllNames:
    """给 checkin.sender_name 用：四张名册 + 支局长名单里的所有姓名。"""

    def __init__(self, people: dict) -> None:
        self.people = people


class DailyRosters:
    """
    国庆每天的活动名单（09-29 用户：活动人员每天不同；我提前做好放后台，到哪天用哪天）。
    data/roster_days.json：{"days": {"2026-10-01": {"城区": {"teams", "people"}, "农村": {...}}, ...}}，
    每天每区的形状和 roster.json 一样。由 deploy/make_roster_days.py 从计划表生成。
    """

    def __init__(self, path: Path) -> None:
        self.path = Path(path)
        self.days: dict[str, dict[str, Roster]] = {}
        self._mtime = None
        self.reload()

    def reload(self) -> None:
        self.days = {}
        try:
            raw = json.loads(self.path.read_text(encoding="utf-8"))
            self._mtime = self.path.stat().st_mtime
        except FileNotFoundError:
            self._mtime = None
            return
        except Exception as exc:
            log.error("每天的活动名单读不了：%s", exc)
            return
        for day, regs in (raw.get("days") or {}).items():
            self.days[day] = {r: Roster.from_dict(v, f"{day} {r}") for r, v in (regs or {}).items() if r in REGIONS}
        log.info("每天的活动名单：%s", "、".join(sorted(self.days)) or "无")

    def maybe_reload(self) -> bool:
        try:
            mtime = self.path.stat().st_mtime
        except OSError:
            mtime = None
        if mtime == self._mtime:
            return False
        self.reload()
        return True

    def get(self, day: str) -> dict | None:
        return self.days.get((day or "")[:10])

    def names(self) -> set[str]:
        return {n for regs in self.days.values() for ros in regs.values() for n in ros.people}


class RosterSet:
    def __init__(self, data_dir: Path = DATA) -> None:
        d = Path(data_dir)
        self.activity = {r: Roster(d / ACTIVITY_FILES[r]) for r in REGIONS}
        self.branch = {r: BranchRoster(d / BRANCH_FILES[r]) for r in REGIONS}
        self.chiefs = ChiefList(d / CHIEFS_FILE)
        self.daily = DailyRosters(d / DAILY_FILE)

    def maybe_reload(self) -> bool:
        changed = False
        for r in REGIONS:
            changed |= self.activity[r].maybe_reload()
            changed |= self.branch[r].maybe_reload()
        changed |= self.chiefs.maybe_reload()
        changed |= self.daily.maybe_reload()
        return changed

    def on(self, day: str | None) -> "RosterSet":
        """
        那一天的名册（国庆每天名单不同，09-29）：有那天的活动名单就用那天的，没有就用平时的
        roster.json / roster_rural.json。支局名册、支局长名单不分天。只读视图，别在上面改名册。
        """
        rosters = self.daily.get(day or "")
        if not rosters:
            return self
        v = copy.copy(self)
        v.activity = {r: rosters.get(r) or self.activity[r] for r in REGIONS}
        return v

    def all_names(self) -> _AllNames:
        people: dict[str, dict] = {}
        for r in REGIONS:
            people.update({n: {} for n in self.branch[r].people})
            people.update({n: {} for n in self.activity[r].people})
        people.update({n: {} for n in self.daily.names()})       # 7 天名单里的人都认得（账号拼音对人）
        people.update({n: {} for n in self.chiefs.names()})
        return _AllNames(people)

    def identify_name(self, name: str, prefer_branch: str = "") -> Identity:
        """
        姓名 → 身份。同一个名字在两个区都有（城区东林、农村安稳各有一个杨鹏，09-29）时，
        按 prefer_branch（晒单写的支局 / 水印里的支局）挑那个支局的；没给或对不上取第一个。
        """
        name = (name or "").strip()
        if not name:
            return Identity()
        hits = []
        for r in REGIONS:
            ros = self.activity[r]
            if name in ros.people:
                grp, team = ros.lookup(name)
                b = self.branch[r].people.get(name) or {}
                info = ros.people.get(name) or {}
                hits.append(Identity(name, r, ACTIVITY, team, grp, b.get("branch") or branch_stem(team),
                                     bool(b.get("community") or info.get("community"))))
        if not hits:
            for r in REGIONS:
                b = self.branch[r].people.get(name)
                if b:
                    hits.append(Identity(name, r, OTHER, "", "", b.get("branch", ""), bool(b.get("community"))))
        if not hits:
            return Identity(name=name)
        want = branch_stem(prefer_branch)
        return next((h for h in hits if want and h.branch == want), hits[0])

    def identify_sender(self, userid: str, prefer_branch: str = "") -> Identity:
        from .checkin import sender_name
        return self.identify_name(sender_name(userid, self.all_names()), prefer_branch)

    def identify_members(self, members: str) -> Identity:
        """打卡第 2 步：水印「小组成员」里出现的活动人员（按出现次数最多的那个组）。"""
        compact = re.sub(r"[\s、,，/]", "", members or "")
        if not compact:
            return Identity()
        best, hits = Identity(), 0
        for r in REGIONS:
            counts: dict[str, int] = {}
            first: dict[str, str] = {}
            for name in self.activity[r].people:
                if name and name in compact:
                    g = self.activity[r].people[name].get("group", "")
                    counts[g] = counts.get(g, 0) + 1
                    first.setdefault(g, name)
            for g, c in counts.items():
                if c > hits:
                    best, hits = self.identify_name(first[g]), c
        return best

    def import_named(self, which: str, data: bytes) -> str:
        """管理员单聊「农村活动名册」+ Excel 这种：导进对应那张。"""
        region, kind = ROSTER_COMMANDS[which]
        if kind == ACTIVITY:
            summary = self.activity[region].import_workbook(data)
        elif kind == "支局":
            summary = self.branch[region].import_workbook(data)
        else:
            summary = self.chiefs.import_workbook(data)
        # 活动名册里的人应该在本区支局名册里 —— 只提醒，不拦
        if kind == ACTIVITY and self.branch[region].people:
            missing = [n for n in self.activity[region].people if n not in self.branch[region].people]
            if missing:
                summary += f"；注意：{len(missing)} 人不在{region}支局名册里（{'、'.join(missing[:5])}{'…' if len(missing) > 5 else ''}）"
        return summary
